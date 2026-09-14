#!/usr/bin/env python3
"""Prepare, validate, and write British-period eight-column asset sheets."""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ASSET_TYPE = "\u8d44\u4ea7\u7c7b\u578b"
CN_NAME = "\u8d44\u4ea7\u4e2d\u6587\u539f\u540d"
EN_NAME = "\u8d44\u4ea7\u672c\u5730\u5316\u82f1\u6587\u540d"
DESC = "\u8d44\u4ea7\u7b80\u4ecb"
ORIGINAL_DESC = "\u539f\u59cb\u4e2d\u6587\u89e3\u6790\u63cf\u8ff0"
CN_STATE_REF = "\u89d2\u8272\u72b6\u6001\u4e2d\u6587\u53c2\u8003\u63cf\u8ff0"
LOCAL_STATE_REF = "\u89d2\u8272\u72b6\u6001\u672c\u5730\u5316\u53c2\u8003"
EPISODE = "\u51fa\u73b0\u96c6\u53f7"
COLUMNS = [ASSET_TYPE, CN_NAME, EN_NAME, DESC, ORIGINAL_DESC, CN_STATE_REF, LOCAL_STATE_REF, EPISODE]

SRC_TYPE = "\u8d44\u4ea7\u7c7b\u578b"
SRC_NAME = "\u8d44\u4ea7\u540d"
SRC_DESC = "\u8d44\u4ea7\u7b80\u4ecb\uff08\u542b\u8be6\u7ec6\u670d\u9970\u63cf\u8ff0\uff1a\u8863\u670d+\u978b\u5b50+\u914d\u9970\uff09"
SRC_PROMPT = "\u9002\u914d\u6587\u751f\u56fe\u7684\u63d0\u793a\u8bcd"
SRC_EPISODES = "\u51fa\u73b0\u96c6\u6570"
SRC_NOTES = "\u5907\u6ce8"

CHARACTER = "\u4eba\u7269"
SCENE = "\u573a\u666f"
PROP = "\u9053\u5177"
VALID_TYPES = {CHARACTER, SCENE, PROP}
DEFAULT_STATE = "\u9ed8\u8ba4\u6001"
SHOT_ASSET_KEYWORDS = [
    "\u8d44\u4ea7",
    "\u4eba\u7269",
    "\u89d2\u8272",
    "\u573a\u666f",
    "\u9053\u5177",
    "\u670d\u5316\u9053",
    "\u6d89\u53ca",
    "\u51fa\u73b0",
    "\u753b\u9762",
    "\u63cf\u8ff0",
    "\u63d0\u793a\u8bcd",
]

DIALOGUE_EP = "\u5267\u96c6"
DIALOGUE_TIME = "\u65f6\u95f4/\u955c\u5934"
DIALOGUE_CN_SPEAKER = "\u4e2d\u6587\u4eba\u7269"
DIALOGUE_LOCAL_SPEAKER = "\u672c\u5730\u5316\u4eba\u7269"
DIALOGUE_CONTEXT = "\u573a\u666f/\u8bed\u5883"
DIALOGUE_CN_LINE = "\u4e2d\u6587\u53f0\u8bcd"
DIALOGUE_LOCAL_LINE = "\u82f1\u56fd\u672c\u5730\u5316\u53f0\u8bcd"
DIALOGUE_NOTES = "\u5907\u6ce8"
DIALOGUE_COLUMNS = [
    DIALOGUE_EP,
    DIALOGUE_TIME,
    DIALOGUE_CN_SPEAKER,
    DIALOGUE_LOCAL_SPEAKER,
    DIALOGUE_CONTEXT,
    DIALOGUE_CN_LINE,
    DIALOGUE_LOCAL_LINE,
    DIALOGUE_NOTES,
]

EPISODE_HEADER_CANDIDATES = ["EP", "\u5267\u96c6", "\u96c6\u6570", "\u96c6\u53f7", "\u51fa\u73b0\u96c6\u53f7"]
TIME_HEADER_CANDIDATES = ["\u65f6\u95f4", "\u8d77\u6b62", "\u65f6\u95f4\u7801", "\u955c\u5934", "\u955c\u5934\u7ec4", "\u5206\u955c\u53f7"]
SPEAKER_HEADER_CANDIDATES = ["\u4eba\u7269", "\u89d2\u8272", "\u8bf4\u8bdd\u4eba", "\u53d1\u8a00\u4eba", "\u4e2d\u6587\u4eba\u7269"]
LINE_HEADER_CANDIDATES = ["\u53f0\u8bcd", "\u4e2d\u6587\u53f0\u8bcd", "\u5bf9\u767d", "\u5b57\u5e55", "\u6587\u672c", "\u53e3\u64ad"]
CONTEXT_HEADER_CANDIDATES = ["\u573a\u666f", "\u8bed\u5883", "\u60c5\u5883", "\u753b\u9762", "\u52a8\u4f5c", "\u955c\u5934\u63cf\u8ff0", "\u5907\u6ce8"]


def load_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key:
            values[key] = value
    return values


def find_input_files(input_dir: Path) -> list[Path]:
    if not input_dir.exists():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")
    files = [path for path in input_dir.glob("*.xlsx") if path.is_file()]
    if not files:
        raise ValueError(f"No .xlsx files found in {input_dir}")
    return sorted(files, key=lambda path: episode_number(path.name))


def episode_number(name: str) -> tuple[int, str]:
    match = re.search(r"EP\s*0*(\d+)", name, flags=re.IGNORECASE)
    if match:
        return (int(match.group(1)), name)
    match = re.search(r"(\d+)", name)
    if match:
        return (int(match.group(1)), name)
    return (999999, name)


def episode_label(path: Path, fallback_index: int) -> str:
    num, _ = episode_number(path.name)
    if num != 999999:
        return f"EP{num}"
    return f"EP{fallback_index}"


def normalize_episode_value(value: str, default_label: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        return default_label
    nums = re.findall(r"\d+", raw)
    if nums:
        return ",".join(f"EP{int(num)}" for num in nums)
    chinese_digits = {
        "\u4e00": 1,
        "\u4e8c": 2,
        "\u4e09": 3,
        "\u56db": 4,
        "\u4e94": 5,
        "\u516d": 6,
        "\u4e03": 7,
        "\u516b": 8,
        "\u4e5d": 9,
        "\u5341": 10,
    }
    labels = [f"EP{value}" for ch, value in chinese_digits.items() if ch in raw]
    return ",".join(labels) if labels else default_label


def episode_sort_key(label: str) -> tuple[int, str]:
    match = re.search(r"EP\s*0*(\d+)", label or "", flags=re.IGNORECASE)
    if match:
        return (int(match.group(1)), label)
    match = re.search(r"\d+", label or "")
    if match:
        return (int(match.group(0)), label)
    return (999999, label or "")


def merge_episode_values(*values: str) -> str:
    labels: list[str] = []
    seen: set[str] = set()
    for value in values:
        normalized = normalize_episode_value(value, "")
        for label in [item.strip() for item in normalized.split(",") if item.strip()]:
            if label not in seen:
                seen.add(label)
                labels.append(label)
    return ",".join(sorted(labels, key=episode_sort_key))


def merge_full_series_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    merged: dict[tuple[str, str], dict[str, str]] = {}
    order: list[tuple[str, str]] = []
    for row in rows:
        key = (row.get(ASSET_TYPE, ""), row.get(CN_NAME, ""))
        if key not in merged:
            merged[key] = dict(row)
            order.append(key)
            continue

        current = merged[key]
        if row.get(EN_NAME) and not current.get(EN_NAME):
            current[EN_NAME] = row.get(EN_NAME, "")
        for field in (DESC, ORIGINAL_DESC, CN_STATE_REF, LOCAL_STATE_REF):
            if len(row.get(field, "")) > len(current.get(field, "")):
                current[field] = row.get(field, "")
        current[EPISODE] = merge_episode_values(current.get(EPISODE, ""), row.get(EPISODE, ""))

    return [merged[key] for key in order]


def ensure_cn_state_separator(name: str, asset_type: str) -> str:
    name = str(name or "").strip()
    if asset_type == CHARACTER:
        name = name.replace("\u00b7", "-")
        if name and "-" not in name and "\uff0d" not in name:
            return f"{name}-{DEFAULT_STATE}"
    return name


def normalize_asset_type(value: str) -> str:
    raw = str(value or "").strip()
    if raw in VALID_TYPES:
        return raw
    for broad_type in (CHARACTER, SCENE, PROP):
        if raw.startswith(broad_type):
            return broad_type
    return PROP


def state_suffix(cn_name: str) -> str:
    normalized = cn_name.replace("\uff0d", "-").replace("\u00b7", "-")
    if "-" in normalized:
        suffix = normalized.rsplit("-", 1)[1].strip()
        return suffix or DEFAULT_STATE
    return DEFAULT_STATE


def character_base_name(cn_name: str) -> str:
    normalized = str(cn_name or "").replace("\uff0d", "-").replace("\u00b7", "-")
    return normalized.rsplit("-", 1)[0].strip() if "-" in normalized else normalized.strip()


def read_source_rows(path: Path, default_ep: str) -> list[dict[str, str]]:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    rows: list[dict[str, str]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        item = {header: str(value or "").strip() for header, value in zip(headers, row)}
        asset_type = normalize_asset_type(item.get(SRC_TYPE) or item.get(ASSET_TYPE))
        name = item.get(SRC_NAME) or item.get(CN_NAME)
        if not asset_type and not name:
            continue
        cn_name = ensure_cn_state_separator(name, asset_type)
        desc_parts = []
        if item.get(SRC_DESC):
            desc_parts.append(item[SRC_DESC])
        if item.get(SRC_PROMPT):
            desc_parts.append(f"\u539f\u59cb\u63d0\u793a\u8bcd\uff1a{item[SRC_PROMPT]}")
        if item.get(SRC_NOTES):
            desc_parts.append(f"\u539f\u59cb\u5907\u6ce8\uff1a{item[SRC_NOTES]}")
        original_desc = item.get(ORIGINAL_DESC) or "\uff1b".join(desc_parts)
        rows.append(
            {
                ASSET_TYPE: asset_type,
                CN_NAME: cn_name,
                EN_NAME: "",
                DESC: "",
                ORIGINAL_DESC: original_desc,
                CN_STATE_REF: item.get(CN_STATE_REF, ""),
                LOCAL_STATE_REF: "",
                EPISODE: normalize_episode_value(item.get(SRC_EPISODES) or item.get(EPISODE), default_ep),
            }
        )
    return rows


def normalize_completed_row(row: dict[str, Any], default_ep: str) -> dict[str, str]:
    asset_type = normalize_asset_type(row.get(ASSET_TYPE, ""))
    cn_name = ensure_cn_state_separator(str(row.get(CN_NAME, "") or "").strip(), asset_type)
    en_name = str(row.get(EN_NAME, "") or "").strip()
    if asset_type == CHARACTER and en_name:
        suffix = state_suffix(cn_name)
        if "-" not in en_name and "\uff0d" not in en_name:
            en_name = f"{en_name}-{suffix}"
    desc = str(row.get(DESC, "") or "").strip()
    original_desc = str(row.get(ORIGINAL_DESC, "") or "").strip()
    cn_state_ref = str(row.get(CN_STATE_REF, "") or "").strip()
    local_state_ref = str(row.get(LOCAL_STATE_REF, "") or "").strip()
    ep = normalize_episode_value(str(row.get(EPISODE, "") or ""), default_ep)
    return {
        ASSET_TYPE: asset_type,
        CN_NAME: cn_name,
        EN_NAME: en_name,
        DESC: desc,
        ORIGINAL_DESC: original_desc,
        CN_STATE_REF: cn_state_ref,
        LOCAL_STATE_REF: local_state_ref,
        EPISODE: ep,
    }


def validate_completed_rows(rows: list[dict[str, str]]) -> list[str]:
    errors: list[str] = []
    character_states: dict[str, set[str]] = {}
    character_bases: set[str] = set()
    for row in rows:
        if row.get(ASSET_TYPE, "") != CHARACTER:
            continue
        base = character_base_name(row.get(CN_NAME, ""))
        character_bases.add(base)
        character_states.setdefault(base, set()).add(state_suffix(row.get(CN_NAME, "")))
    for index, row in enumerate(rows, start=2):
        if list(row.keys()) != COLUMNS:
            errors.append(f"row {index}: columns must exactly match required eight-column schema")
        asset_type = row.get(ASSET_TYPE, "")
        cn_name = row.get(CN_NAME, "")
        en_name = row.get(EN_NAME, "")
        desc = row.get(DESC, "")
        local_state_ref = row.get(LOCAL_STATE_REF, "")
        episode = row.get(EPISODE, "")
        if not desc:
            errors.append(f"row {index}: localized asset brief is required")
        if asset_type == CHARACTER:
            if "-" not in cn_name and "\uff0d" not in cn_name:
                errors.append(f"row {index}: character Chinese name must include state suffix with hyphen")
            if not en_name:
                errors.append(f"row {index}: character localized English name is required")
            elif "-" not in en_name and "\uff0d" not in en_name:
                errors.append(f"row {index}: character localized English name must include Chinese state suffix")
            if not local_state_ref:
                errors.append(f"row {index}: character localized state reference is required")
            base = character_base_name(cn_name)
            own_state = state_suffix(cn_name)
            sibling_states = character_states.get(base, set()) - {own_state}
            contaminated = sorted(state for state in sibling_states if state and (state in desc or state in local_state_ref))
            if contaminated:
                errors.append(
                    f"row {index}: character state prompt contains sibling states: {', '.join(contaminated[:6])}"
                )
            original_desc = row.get(ORIGINAL_DESC, "")
            subject_match = re.match(r"^([\u4e00-\u9fff]{2,8})\u5728", original_desc)
            leading_identity = subject_match.group(1) if subject_match else ""
            if leading_identity not in character_bases or leading_identity == base or base in original_desc:
                leading_identity = ""
            if leading_identity:
                errors.append(
                    f"row {index}: source state description starts with a different character identity: {leading_identity}"
                )
        elif not en_name:
            errors.append(f"row {index}: scene/prop localized English name is required")
        if re.search(r"EP[\u4e00-\u9fff]+", episode, flags=re.IGNORECASE):
            errors.append(f"row {index}: episode value must not mix EP with Chinese numerals")
    return errors


def write_excel(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "\u82f1\u56fd\u672c\u5730\u5316\u8d44\u4ea7\u8868"
    sheet.append(COLUMNS)
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    type_order = {CHARACTER: 0, SCENE: 1, PROP: 2}
    for row in sorted(rows, key=lambda r: (type_order.get(r.get(ASSET_TYPE, ""), 99), r.get(CN_NAME, ""))):
        sheet.append([row.get(column, "") for column in COLUMNS])
    for index, width in enumerate([14, 30, 34, 54, 62, 54, 62, 18], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    workbook.save(path)


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def find_dialogue_files(input_path: Path) -> list[Path]:
    if not input_path.exists():
        raise FileNotFoundError(f"Dialogue input not found: {input_path}")
    if input_path.is_file():
        return [input_path]
    files = [
        path
        for pattern in ("*.xlsx", "*.csv")
        for path in input_path.glob(pattern)
        if path.is_file()
    ]
    if not files:
        raise ValueError(f"No .xlsx or .csv dialogue files found in {input_path}")
    return sorted(files, key=lambda path: episode_number(path.name))


def read_tabular_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    return [
                        {str(key or "").strip(): str(value or "").strip() for key, value in row.items()}
                        for row in csv.DictReader(handle)
                    ]
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError("csv", b"", 0, 1, f"Unable to decode {path}")
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    rows: list[dict[str, str]] = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        item = {header: str(value or "").strip() for header, value in zip(headers, raw_row)}
        if any(item.values()):
            rows.append(item)
    return rows


def pick_value(row: dict[str, str], candidates: list[str]) -> str:
    for candidate in candidates:
        if candidate in row and row[candidate]:
            return row[candidate]
    for key, value in row.items():
        if any(candidate in key for candidate in candidates) and value:
            return value
    return ""


def build_character_context(asset_rows: list[dict[str, str]]) -> tuple[list[dict[str, str]], dict[str, str]]:
    context: list[dict[str, str]] = []
    speaker_map: dict[str, str] = {}
    for row in asset_rows:
        if row.get(ASSET_TYPE) != CHARACTER:
            continue
        cn_name = row.get(CN_NAME, "")
        en_name = row.get(EN_NAME, "")
        if not cn_name or not en_name:
            continue
        base, _ = split_cn_name(cn_name)
        context.append(
            {
                DIALOGUE_CN_SPEAKER: cn_name,
                DIALOGUE_LOCAL_SPEAKER: en_name,
                DESC: row.get(DESC, ""),
                EPISODE: row.get(EPISODE, ""),
            }
        )
        speaker_map[cn_name] = en_name
        if base:
            speaker_map.setdefault(base, split_cn_name(en_name)[0] or en_name)
    return context, dict(sorted(speaker_map.items(), key=lambda item: len(item[0]), reverse=True))


def match_local_speaker(speaker: str, speaker_map: dict[str, str]) -> str:
    speaker = str(speaker or "").strip()
    if not speaker:
        return ""
    for source, target in speaker_map.items():
        if source == speaker or source in speaker or speaker in source:
            return target
    return ""


def normalize_dialogue_row(row: dict[str, Any], default_ep: str, speaker_map: dict[str, str] | None = None) -> dict[str, str]:
    speaker_map = speaker_map or {}
    ep = normalize_episode_value(str(row.get(DIALOGUE_EP, "") or ""), default_ep).split(",", 1)[0]
    cn_speaker = str(row.get(DIALOGUE_CN_SPEAKER, "") or "").strip()
    local_speaker = str(row.get(DIALOGUE_LOCAL_SPEAKER, "") or "").strip() or match_local_speaker(cn_speaker, speaker_map)
    return {
        DIALOGUE_EP: ep,
        DIALOGUE_TIME: str(row.get(DIALOGUE_TIME, "") or "").strip(),
        DIALOGUE_CN_SPEAKER: cn_speaker,
        DIALOGUE_LOCAL_SPEAKER: local_speaker,
        DIALOGUE_CONTEXT: str(row.get(DIALOGUE_CONTEXT, "") or "").strip(),
        DIALOGUE_CN_LINE: str(row.get(DIALOGUE_CN_LINE, "") or "").strip(),
        DIALOGUE_LOCAL_LINE: str(row.get(DIALOGUE_LOCAL_LINE, "") or "").strip(),
        DIALOGUE_NOTES: str(row.get(DIALOGUE_NOTES, "") or "").strip(),
    }


def source_dialogue_to_rows(path: Path, default_ep: str, speaker_map: dict[str, str]) -> list[dict[str, str]]:
    rows = []
    for row in read_tabular_rows(path):
        line = pick_value(row, LINE_HEADER_CANDIDATES)
        speaker = pick_value(row, SPEAKER_HEADER_CANDIDATES)
        if not line and not speaker:
            continue
        ep = normalize_episode_value(pick_value(row, EPISODE_HEADER_CANDIDATES), default_ep).split(",", 1)[0]
        rows.append(
            normalize_dialogue_row(
                {
                    DIALOGUE_EP: ep,
                    DIALOGUE_TIME: pick_value(row, TIME_HEADER_CANDIDATES),
                    DIALOGUE_CN_SPEAKER: speaker,
                    DIALOGUE_LOCAL_SPEAKER: match_local_speaker(speaker, speaker_map),
                    DIALOGUE_CONTEXT: pick_value(row, CONTEXT_HEADER_CANDIDATES),
                    DIALOGUE_CN_LINE: line,
                    DIALOGUE_LOCAL_LINE: "",
                    DIALOGUE_NOTES: "",
                },
                ep,
                speaker_map,
            )
        )
    return rows


def write_dialogue_excel(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "\u82f1\u56fd\u672c\u5730\u5316\u53f0\u8bcd\u5bf9\u7167"
    sheet.append(DIALOGUE_COLUMNS)
    fill = PatternFill("solid", fgColor="E8F3E8")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        sheet.append([row.get(column, "") for column in DIALOGUE_COLUMNS])
    for index, width in enumerate([12, 18, 18, 24, 42, 50, 60, 34], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    workbook.save(path)


def read_reference_asset_rows(path: Path, limit: int = 120) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Reference asset file not found: {path}")
    rows: list[dict[str, str]] = []
    for index, row in enumerate(read_tabular_rows(path), start=1):
        if index > limit:
            break
        if not row.get(CN_NAME) and not row.get(ASSET_TYPE):
            continue
        rows.append(normalize_completed_row(row, row.get(EPISODE) or "EP1"))
    return merge_full_series_rows(rows)


def prepare_dialogue(
    env: dict[str, str],
    dialogue_input: Path,
    asset_mapping_file: Path,
    dialogue_draft_dir: Path,
) -> None:
    asset_rows = read_localized_asset_rows(asset_mapping_file)
    character_context, speaker_map = build_character_context(asset_rows)
    files = find_dialogue_files(dialogue_input)
    grouped: dict[str, list[dict[str, str]]] = {}
    manifest = []
    for index, source_file in enumerate(files, start=1):
        default_ep = episode_label(source_file, index)
        for row in source_dialogue_to_rows(source_file, default_ep, speaker_map):
            grouped.setdefault(row[DIALOGUE_EP], []).append(row)
    for ep in sorted(grouped, key=lambda label: episode_number(label)[0]):
        payload = {
            "episode": ep,
            "setting": env.get("BRITISH_SETTING", "fictional Victorian-inspired British kingdom and aristocratic estate society"),
            "asset_mapping_file": str(asset_mapping_file),
            "character_context": character_context,
            "rules_summary": [
                "Codex translates and adapts dialogue itself; do not call Gemini.",
                "Use localized character names from the asset mapping workbook.",
                "Keep original Chinese line and write the British localized counterpart.",
                "Translation must fit period setting, social rank, current environment, relationship, and episode context.",
                "Preserve plot intent, conflict, reveals, emotional temperature, and address hierarchy.",
                "Use duration-friendly equal-meaning rewrites: keep the intent, but shorten English syntax instead of translating Chinese wording literally.",
                "Keep English close to the source speaking time; prefer short names and titles, remove redundant connectors, and aim for no more than 1.2-1.4x source length unless clarity requires it.",
                "For voice-over, compress exposition into dense natural English sentences; for spoken lines, prioritize mouth-friendly concise wording.",
            ],
            "rows_to_translate": grouped[ep],
        }
        draft_json = dialogue_draft_dir / f"{ep}_dialogue_localization_draft.json"
        write_json(draft_json, payload)
        manifest.append({"episode": ep, "draft_json": str(draft_json), "row_count": len(grouped[ep])})
    write_json(dialogue_draft_dir / "dialogue_localization_manifest.json", manifest)
    print(f"Prepared {len(manifest)} dialogue draft JSON files in {dialogue_draft_dir}")


def validate_dialogue_rows(rows: list[dict[str, str]]) -> list[str]:
    errors: list[str] = []
    for index, row in enumerate(rows, start=2):
        if list(row.keys()) != DIALOGUE_COLUMNS:
            errors.append(f"row {index}: columns must exactly match dialogue schema")
        if not row.get(DIALOGUE_CN_LINE):
            errors.append(f"row {index}: Chinese dialogue is required")
        if not row.get(DIALOGUE_LOCAL_LINE):
            errors.append(f"row {index}: British localized dialogue is required")
        if re.search(r"EP[\u4e00-\u9fff]+", row.get(DIALOGUE_EP, ""), flags=re.IGNORECASE):
            errors.append(f"row {index}: episode value must not mix EP with Chinese numerals")
    return errors


def write_dialogue_outputs(dialogue_completed_dir: Path, dialogue_output_dir: Path, full_dialogue_file: Path) -> None:
    files = sorted(dialogue_completed_dir.glob("*_dialogue_completed.json"), key=lambda path: episode_number(path.name))
    if not files:
        raise FileNotFoundError(f"No *_dialogue_completed.json files found in {dialogue_completed_dir}")
    all_rows: list[dict[str, str]] = []
    for index, path in enumerate(files, start=1):
        data = read_json(path)
        ep = str(data.get("episode") or f"EP{index}")
        raw_rows = data.get("completed_rows") or data.get("rows") or []
        rows = [normalize_dialogue_row(row, ep) for row in raw_rows]
        errors = validate_dialogue_rows(rows)
        if errors:
            raise ValueError(f"Dialogue validation failed for {path}:\n" + "\n".join(errors))
        output_file = dialogue_output_dir / f"{ep}_\u82f1\u56fd\u672c\u5730\u5316\u53f0\u8bcd\u5bf9\u7167.xlsx"
        write_dialogue_excel(output_file, rows)
        all_rows.extend(rows)
        print(f"wrote {len(rows)} dialogue rows -> {output_file}")
    write_dialogue_excel(full_dialogue_file, all_rows)
    print(f"wrote full dialogue workbook -> {full_dialogue_file}")


def prepare(env: dict[str, str], input_dir: Path, draft_dir: Path) -> None:
    files = find_input_files(input_dir)
    manifest = []
    previous_rows: list[dict[str, str]] = []
    reference_file = env.get("BRITISH_REFERENCE_ASSET_FILE", "").strip()
    reference_rows: list[dict[str, str]] = []
    if reference_file:
        reference_rows = read_reference_asset_rows(Path(reference_file).expanduser())
    for index, source_file in enumerate(files, start=1):
        ep = episode_label(source_file, index)
        rows = read_source_rows(source_file, ep)
        payload = {
            "episode": ep,
            "source_file": str(source_file),
            "reference_asset_file": reference_file,
            "reference_asset_rows": reference_rows,
            "setting": env.get("BRITISH_SETTING", "fictional Victorian-inspired British kingdom and aristocratic estate society"),
            "rules_summary": [
                "Codex must localize rows itself; do not call Gemini.",
                "If reference_asset_rows is present, follow its naming style, setting vocabulary, prompt structure, costume detail level, and episode formatting unless it conflicts with the current source row.",
                "Keep exactly eight columns in the required order.",
                "Fill localized English name and localized asset brief for every row.",
                "Same-family surnames stay consistent.",
                "No pinyin and no real-world place names.",
                "Character localized English names must end with Chinese state suffix.",
                "Preserve original Chinese description and Chinese state reference verbatim when present.",
                "Character rows must fill localized state reference; scene and prop rows leave both state-reference columns blank unless the source explicitly supplies state evidence.",
                "Scene and prop asset briefs must be British-period localized and include concise image-generation guidance.",
            ],
            "previous_localized_rows": previous_rows,
            "rows_to_localize": rows,
        }
        draft_json = draft_dir / f"{ep}_codex_localization_draft.json"
        write_json(draft_json, payload)
        output_file = Path(env.get("BRITISH_ASSET_OUTPUT_DIR", "british_localized_assets")).expanduser() / f"{ep}_英国本地化资产表.xlsx"
        manifest.append({"episode": ep, "draft_json": str(draft_json), "output_file": str(output_file)})
        previous_rows = rows
    write_json(draft_dir / "localization_manifest.json", manifest)
    print(f"Prepared {len(manifest)} draft JSON files in {draft_dir}")
    print("Codex should edit each draft JSON by adding completed_rows, then run --mode write.")


def write_outputs(completed_dir: Path, output_dir: Path, full_series_file: Path) -> None:
    files = sorted(completed_dir.glob("*_completed.json"), key=lambda path: episode_number(path.name))
    if not files:
        raise FileNotFoundError(f"No *_completed.json files found in {completed_dir}")
    all_rows: list[dict[str, str]] = []
    for index, path in enumerate(files, start=1):
        data = read_json(path)
        ep = str(data.get("episode") or f"EP{index}")
        raw_rows = data.get("completed_rows") or data.get("rows") or []
        rows = [normalize_completed_row(row, ep) for row in raw_rows]
        errors = validate_completed_rows(rows)
        if errors:
            raise ValueError(f"Validation failed for {path}:\n" + "\n".join(errors))
        output_file = output_dir / f"{ep}_英国本地化资产表.xlsx"
        write_excel(output_file, rows)
        all_rows.extend(rows)
        print(f"wrote {len(rows)} rows -> {output_file}")
    merged_rows = merge_full_series_rows(all_rows)
    write_excel(full_series_file, merged_rows)
    print(f"wrote full series workbook -> {full_series_file} ({len(all_rows)} rows merged to {len(merged_rows)} unique assets)")


def read_localized_asset_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Localized asset workbook not found: {path}")
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    rows: list[dict[str, str]] = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        row = {header: str(value or "").strip() for header, value in zip(headers, raw_row)}
        if row.get(CN_NAME):
            rows.append(normalize_completed_row(row, row.get(EPISODE) or "EP1"))
    return rows


def split_cn_name(name: str) -> tuple[str, str]:
    normalized = str(name or "").strip().replace("\uff0d", "-")
    if "-" not in normalized:
        return normalized, ""
    base, suffix = normalized.rsplit("-", 1)
    return base.strip(), suffix.strip()


def localized_display_name(row: dict[str, str]) -> str:
    cn_name = row.get(CN_NAME, "")
    en_name = row.get(EN_NAME, "")
    return en_name or cn_name


def build_replacement_map(rows: list[dict[str, str]]) -> dict[str, str]:
    replacements: dict[str, str] = {}
    for row in rows:
        cn_name = row.get(CN_NAME, "")
        if not cn_name:
            continue
        display = localized_display_name(row)
        if display and display != cn_name:
            replacements[cn_name] = display
        base, suffix = split_cn_name(cn_name)
        if base and display and row.get(ASSET_TYPE) == CHARACTER:
            display_base, _ = split_cn_name(display)
            replacements.setdefault(base, display_base or display)
        if suffix and display and row.get(ASSET_TYPE) == CHARACTER:
            replacements.setdefault(f"{base}{suffix}", display)
    return dict(sorted(replacements.items(), key=lambda item: len(item[0]), reverse=True))


def should_replace_column(header: str, replace_all: bool) -> bool:
    if replace_all:
        return True
    header = str(header or "")
    return any(keyword in header for keyword in SHOT_ASSET_KEYWORDS)


def replace_text(value: Any, replacements: dict[str, str]) -> tuple[Any, list[str]]:
    if value is None or not isinstance(value, str):
        return value, []
    updated = value
    hits: list[str] = []
    for source, target in replacements.items():
        if source and target and source in updated:
            updated = updated.replace(source, target)
            hits.append(source)
    return updated, hits


def shot_group_key(value: Any) -> str:
    match = re.search(r"EP\d+[-_]?G\d+", str(value or ""), flags=re.IGNORECASE)
    if not match:
        return ""
    return match.group(0).replace("_", "-").upper()


def dialogue_speaker_name(value: str) -> str:
    name = str(value or "").strip()
    if "-" in name:
        return name.split("-", 1)[0].strip()
    return name


def build_shot_dialogue_map(dialogue_file: Path) -> dict[str, str]:
    if not dialogue_file.exists():
        return {}
    grouped: dict[str, list[str]] = {}
    for row in read_tabular_rows(dialogue_file):
        key = shot_group_key(row.get(DIALOGUE_TIME, ""))
        line = str(row.get(DIALOGUE_LOCAL_LINE, "")).strip()
        if not key or not line:
            continue
        speaker = dialogue_speaker_name(row.get(DIALOGUE_LOCAL_SPEAKER, "") or row.get(DIALOGUE_CN_SPEAKER, ""))
        context = str(row.get(DIALOGUE_CONTEXT, ""))
        label = "voice-over" if "\u65c1\u767d" in context else ""
        prefix = f"{speaker} {label}".strip() if speaker else label
        grouped.setdefault(key, []).append(f'{prefix}: "{line}"' if prefix else f'"{line}"')
    return {key: "\n".join(lines) for key, lines in grouped.items()}


def strip_embedded_dialogue(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    updated = re.sub(r"[^。\n；;]*?(?:\u65c1\u767d|\u53f0\u8bcd)\uff1a\u201c[^\u201d]*\u201d[。\n；;]*", "", value)
    updated = re.sub(r"\s{2,}", " ", updated).strip()
    updated = re.sub(r"(。){2,}", "。", updated)
    return updated


def map_shot_workbook(
    source_shot_file: Path,
    asset_mapping_file: Path,
    output_file: Path,
    report_file: Path,
    replace_all_columns: bool = False,
    dialogue_file: Path | None = None,
) -> None:
    if not source_shot_file.exists():
        raise FileNotFoundError(f"Shot workbook not found: {source_shot_file}")
    rows = read_localized_asset_rows(asset_mapping_file)
    replacements = build_replacement_map(rows)
    if not replacements:
        raise ValueError(f"No replacements built from {asset_mapping_file}")

    workbook = load_workbook(source_shot_file)
    replacement_counts = {key: 0 for key in replacements}
    touched_cells: list[dict[str, Any]] = []
    dialogue_map = build_shot_dialogue_map(dialogue_file) if dialogue_file else {}
    dialogue_updates = 0
    stripped_dialogue_cells = 0
    for sheet in workbook.worksheets:
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        replace_indexes = {
            index
            for index, header in enumerate(headers, start=1)
            if should_replace_column(header, replace_all_columns)
        }
        if not replace_indexes:
            replace_indexes = set(range(1, sheet.max_column + 1)) if replace_all_columns else set()
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column not in replace_indexes:
                    continue
                old_value = cell.value
                new_value, hits = replace_text(old_value, replacements)
                if hits and new_value != old_value:
                    cell.value = new_value
                    for hit in hits:
                        replacement_counts[hit] += 1
                    touched_cells.append(
                        {
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "sources": hits,
                            "old": old_value,
                            "new": new_value,
                        }
                    )
            if dialogue_map:
                row_values = [cell.value for cell in row]
                shot_cell = None
                dialogue_cell = None
                visual_cell = None
                for index, header in enumerate(headers):
                    if header in {"\u955c\u5934\u7ec4\u5e8f\u53f7", "\u955c\u5934", "\u5206\u955c\u53f7"}:
                        shot_cell = row[index]
                    elif header == "\u53f0\u8bcd/\u58f0\u97f3":
                        dialogue_cell = row[index]
                    elif header in {"\u753b\u9762\u5185\u5bb9", "\u955c\u5934\u63cf\u8ff0"}:
                        visual_cell = row[index]
                key = shot_group_key(shot_cell.value if shot_cell else "")
                if key and dialogue_cell and key in dialogue_map and dialogue_cell.value != dialogue_map[key]:
                    dialogue_cell.value = dialogue_map[key]
                    dialogue_updates += 1
                if visual_cell:
                    old_visual = visual_cell.value
                    new_visual = strip_embedded_dialogue(old_visual)
                    if new_visual != old_visual:
                        visual_cell.value = new_visual
                        stripped_dialogue_cells += 1

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_file)
    used = {key: count for key, count in replacement_counts.items() if count}
    unused = [key for key, count in replacement_counts.items() if not count]
    report = {
        "source_shot_file": str(source_shot_file),
        "asset_mapping_file": str(asset_mapping_file),
        "output_file": str(output_file),
        "replace_all_columns": replace_all_columns,
        "replacement_count": sum(used.values()),
        "dialogue_file": str(dialogue_file) if dialogue_file else "",
        "dialogue_updates": dialogue_updates,
        "stripped_dialogue_cells": stripped_dialogue_cells,
        "used_replacements": used,
        "unused_replacements": unused,
        "touched_cells": touched_cells,
    }
    write_json(report_file, report)
    print(f"wrote localized shot workbook -> {output_file}")
    print(f"wrote replacement report -> {report_file}")
    print(f"replaced {sum(used.values())} asset mentions")
    if dialogue_map:
        print(f"updated {dialogue_updates} shot dialogue cells")
        print(f"removed embedded source dialogue from {stripped_dialogue_cells} visual cells")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Prepare/write British-period eight-column asset sheets without Gemini calls.")
    parser.add_argument("--env-file", default=".env", help="Path to .env config.")
    parser.add_argument("--mode", choices=["prepare", "write", "prepare-dialogue", "write-dialogue", "map-shots", "dry-run"], default="prepare")
    parser.add_argument("--input-dir", help="Input directory containing source asset .xlsx files.")
    parser.add_argument("--draft-dir", help="Directory for Codex-editable draft JSON files.")
    parser.add_argument("--completed-dir", help="Directory containing *_completed.json files.")
    parser.add_argument("--output-dir", help="Output directory for localized .xlsx files.")
    parser.add_argument("--shot-file", help="Source whole-series shot group workbook.")
    parser.add_argument("--asset-map-file", help="Localized full-series asset mapping workbook.")
    parser.add_argument("--shot-output-file", help="Output localized shot group workbook.")
    parser.add_argument("--shot-report-file", help="Output JSON replacement report.")
    parser.add_argument("--replace-all-shot-columns", action="store_true", help="Replace asset names in every shot workbook column.")
    parser.add_argument("--dialogue-input", help="Source dialogue table file or directory.")
    parser.add_argument("--dialogue-draft-dir", help="Directory for Codex-editable dialogue draft JSON files.")
    parser.add_argument("--dialogue-completed-dir", help="Directory containing *_dialogue_completed.json files.")
    parser.add_argument("--dialogue-output-dir", help="Output directory for localized dialogue workbooks.")
    parser.add_argument("--dialogue-full-file", help="Output full-series localized dialogue workbook.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    env = load_env(Path(args.env_file))
    input_dir = Path(args.input_dir or env.get("BRITISH_ASSET_INPUT_DIR", "gemini_asset_excels")).expanduser()
    draft_dir = Path(args.draft_dir or env.get("BRITISH_ASSET_DRAFT_DIR", "british_localized_assets/drafts")).expanduser()
    completed_dir = Path(args.completed_dir or env.get("BRITISH_ASSET_COMPLETED_DIR", str(draft_dir))).expanduser()
    output_dir = Path(args.output_dir or env.get("BRITISH_ASSET_OUTPUT_DIR", "british_localized_assets")).expanduser()
    full_series_file = Path(env.get("BRITISH_ASSET_FULL_SERIES_FILE", str(output_dir / "\u5168\u5267_\u82f1\u56fd\u672c\u5730\u5316\u8d44\u4ea7\u8868.xlsx"))).expanduser()
    shot_file = Path(args.shot_file or env.get("BRITISH_SHOT_INPUT_FILE", "gemini_shot_group_excels/\u5168\u5267_\u955c\u5934\u7ec4.xlsx")).expanduser()
    asset_map_file = Path(args.asset_map_file or env.get("BRITISH_SHOT_ASSET_MAP_FILE", str(full_series_file))).expanduser()
    shot_output_file = Path(args.shot_output_file or env.get("BRITISH_SHOT_OUTPUT_FILE", str(output_dir / "\u5168\u5267_\u82f1\u56fd\u672c\u5730\u5316\u955c\u5934\u7ec4.xlsx"))).expanduser()
    shot_report_file = Path(args.shot_report_file or env.get("BRITISH_SHOT_REPORT_FILE", str(output_dir / "\u5168\u5267_\u955c\u5934\u7ec4\u8d44\u4ea7\u66ff\u6362\u62a5\u544a.json"))).expanduser()
    dialogue_input = Path(args.dialogue_input or env.get("BRITISH_DIALOGUE_INPUT", "dialogue_tables")).expanduser()
    dialogue_draft_dir = Path(args.dialogue_draft_dir or env.get("BRITISH_DIALOGUE_DRAFT_DIR", str(output_dir / "dialogue_drafts"))).expanduser()
    dialogue_completed_dir = Path(args.dialogue_completed_dir or env.get("BRITISH_DIALOGUE_COMPLETED_DIR", str(dialogue_draft_dir))).expanduser()
    dialogue_output_dir = Path(args.dialogue_output_dir or env.get("BRITISH_DIALOGUE_OUTPUT_DIR", str(output_dir / "dialogue"))).expanduser()
    dialogue_full_file = Path(args.dialogue_full_file or env.get("BRITISH_DIALOGUE_FULL_FILE", str(dialogue_output_dir / "\u5168\u5267_\u82f1\u56fd\u672c\u5730\u5316\u53f0\u8bcd\u5bf9\u7167.xlsx"))).expanduser()

    if args.mode == "dry-run":
        if not input_dir.exists():
            print(f"Input directory not found yet: {input_dir}")
            print("Run the previous asset inventory skill first, then rerun this skill.")
            return 0
        for index, path in enumerate(find_input_files(input_dir), start=1):
            ep = episode_label(path, index)
            print(f"{ep}: {path} -> draft {draft_dir / f'{ep}_codex_localization_draft.json'}")
        print(f"Output directory -> {output_dir}")
        print(f"Full series -> {full_series_file}")
        print(f"Shot input -> {shot_file}")
        print(f"Shot asset map -> {asset_map_file}")
        print(f"Shot output -> {shot_output_file}")
        print(f"Dialogue input -> {dialogue_input}")
        print(f"Dialogue drafts -> {dialogue_draft_dir}")
        print(f"Dialogue output -> {dialogue_output_dir}")
        print(f"Full dialogue -> {dialogue_full_file}")
        return 0
    if args.mode == "prepare":
        prepare(env, input_dir, draft_dir)
        return 0
    if args.mode == "prepare-dialogue":
        prepare_dialogue(
            env=env,
            dialogue_input=dialogue_input,
            asset_mapping_file=asset_map_file,
            dialogue_draft_dir=dialogue_draft_dir,
        )
        return 0
    if args.mode == "write-dialogue":
        write_dialogue_outputs(
            dialogue_completed_dir=dialogue_completed_dir,
            dialogue_output_dir=dialogue_output_dir,
            full_dialogue_file=dialogue_full_file,
        )
        return 0
    if args.mode == "map-shots":
        map_shot_workbook(
            source_shot_file=shot_file,
            asset_mapping_file=asset_map_file,
            output_file=shot_output_file,
            report_file=shot_report_file,
            replace_all_columns=args.replace_all_shot_columns
            or env.get("BRITISH_SHOT_REPLACE_ALL_COLUMNS", "").lower() in {"1", "true", "yes"},
            dialogue_file=dialogue_full_file,
        )
        return 0
    write_outputs(completed_dir, output_dir, full_series_file)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
