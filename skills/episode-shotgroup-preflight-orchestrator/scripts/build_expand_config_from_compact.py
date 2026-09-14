#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def text_parts(*values: object) -> list[str]:
    parts: list[str] = []
    for value in values:
        text = clean(value)
        if text and text not in parts:
            parts.append(text)
    return parts


def has_cjk(text: str) -> bool:
    return bool(re.search(r"[\u3400-\u9fff]", text or ""))


def read_asset_rules(asset_workbook: Path) -> tuple[list[dict], list[dict], list[dict], str]:
    wb = openpyxl.load_workbook(asset_workbook, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = [clean(v) for v in rows[0]][:8] if rows else []
    expected = ["资产类型", "资产中文原名", "资产本地化英文名", "资产简介", "原始中文解析描述", "角色状态中文参考描述", "角色状态本地化参考", "出现集号"]
    if headers != expected:
        raise RuntimeError(f"asset workbook headers do not match eight-column contract: {headers}")
    people: list[dict] = []
    scenes: list[dict] = []
    props: list[dict] = []
    default_scene = "未入库场景（候选新增/需复核）"
    for row in rows[1:]:
        kind, chinese, localized, intro, source_desc, state_cn, state_local, episodes = [clean(v) for v in row[:8]]
        aliases = text_parts(chinese, localized, intro, source_desc, state_cn, state_local)
        rule = {
            "asset": localized or chinese,
            "aliases": aliases,
            "keywords": aliases,
            "start": 0,
            "end": 9999,
        }
        if kind == "人物" and rule["asset"]:
            people.append(rule)
        elif kind == "场景" and rule["asset"]:
            scenes.append(rule)
            if default_scene.startswith("未入库"):
                default_scene = rule["asset"]
        elif kind == "道具" and rule["asset"]:
            props.append(rule)
    wb.close()
    return people, scenes, props, default_scene


def alias_score(blob: str, rule: dict) -> int:
    score = 0
    for alias in rule.get("aliases") or []:
        text = clean(alias)
        if not text:
            continue
        if has_cjk(text) and text in blob:
            score += 4 + min(4, len(text))
        elif not has_cjk(text) and text.lower() in blob.lower():
            score += 3
        else:
            for token in re.split(r"[\s,;；，。/、()（）]+", text):
                token = token.strip()
                if len(token) >= 2 and token in blob:
                    score += 1
    return score


def record_scene(record: dict, scenes: list[dict], default_scene: str) -> str:
    blob = " ".join(clean(record.get(key)) for key in ("fixed_scene_evidence", "scene_observation", "scene_location", "scene", "beat_summary", "primary_visible_subject"))
    scored = sorted(((alias_score(blob, rule), rule) for rule in scenes), key=lambda item: item[0], reverse=True)
    if scored and scored[0][0] > 0:
        return scored[0][1]["asset"]
    for key in ("fixed_scene_evidence", "scene_observation", "scene_location", "scene"):
        value = clean(record.get(key))
        if value and value.lower() not in {"none", "null", "unclear", "unknown"}:
            return default_scene
    return default_scene


def record_summary(record: dict) -> str:
    return clean(record.get("visible_action") or record.get("visual_action") or record.get("narrative_summary") or "当前镜头组")


def build_groups(records: list[dict], scenes: list[dict], default_scene: str) -> list[dict]:
    groups: list[dict] = []
    sorted_records = sorted(records, key=lambda item: (float(item.get("start_seconds", 0)), float(item.get("end_seconds", 0))))
    buckets: list[list[dict]] = []
    current_key: object = object()
    for record in sorted_records:
        key = record.get("narrative_group")
        if not buckets or key != current_key:
            buckets.append([])
            current_key = key
        buckets[-1].append(record)
    for bucket in buckets:
        start = min(float(record["start_seconds"]) for record in bucket)
        end = max(float(record["end_seconds"]) for record in bucket)
        scene_votes: dict[str, int] = {}
        for record in bucket:
            scene_name = record_scene(record, scenes, default_scene)
            scene_votes[scene_name] = scene_votes.get(scene_name, 0) + 1
        scene = sorted(scene_votes.items(), key=lambda item: item[1], reverse=True)[0][0]
        summaries = [record_summary(record) for record in bucket]
        summary = "；".join(item for item in summaries if item) or "当前镜头组"
        tod_values = [clean(record.get("time_of_day")) for record in bucket if clean(record.get("time_of_day"))]
        time_of_day = next((value for value in tod_values if value != "unclear"), "unclear")
        tod_evidence = next((clean(record.get("time_of_day_evidence")) for record in bucket if clean(record.get("time_of_day_evidence"))), "由聚合 compact record 的昼夜证据继承；不足则标记 unclear。")
        groups.append(
            {
                "start": round(start, 3),
                "end": round(end, 3),
                "scene": scene,
                "summary": summary,
                "time_of_day": time_of_day,
                "time_of_day_evidence": tod_evidence,
            }
        )
    return groups


def read_external_translations(path: Path | None) -> dict[str, str]:
    if not path:
        return {}
    payload = load_json(path)
    if isinstance(payload, dict):
        if isinstance(payload.get("translations"), dict):
            return {clean(k): clean(v) for k, v in payload["translations"].items() if clean(k) and clean(v)}
        return {clean(k): clean(v) for k, v in payload.items() if clean(k) and clean(v)}
    return {}


def build_dialogue_translations(compact: dict, external: dict[str, str]) -> tuple[dict[str, str], dict[str, str], list[str]]:
    translations = {
        "你老公出轨了我老婆": "He's with my wife.",
        "佳佳不是我闺女": "Jiajia is not my daughter.",
        "这怎么可能": "How can that be possible?",
        "老婆，你这衣服底下怎么有条领带啊": "Darling, why is there a tie under your clothes?",
        "老公，这是我集体游戏用的道具，我正在找呢": "Darling, it is just a prop from a group game. I have been looking for it.",
        "谢谢老公": "Thank you, darling.",
        "难道她": "Could she be...",
        "喂老婆": "Hello, darling.",
        "你在哪呢": "Where are you?",
        "公司加班呢": "Working late at office.",
        "欺人太甚": "This is beyond insulting.",
    }
    translations.update(external)
    by_id: dict[str, str] = {}
    warnings: list[str] = []
    for event in compact.get("dialogue_events") or []:
        if not isinstance(event, dict):
            continue
        dialogue_id = clean(event.get("dialogue_id"))
        source_text = clean(event.get("text"))
        if dialogue_id and source_text:
            translated = translations.get(source_text)
            if not translated:
                translated = f"Untranslated dialogue {dialogue_id}; localise before generation."
                warnings.append(f"{dialogue_id}: missing British-localized translation for source dialogue")
            by_id[dialogue_id] = translated
    return translations, by_id, warnings


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--compact", required=True, type=Path)
    parser.add_argument("--assets", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--translations-json", type=Path, help="Optional JSON mapping Chinese dialogue text to British-localized English.")
    args = parser.parse_args()
    compact = load_json(args.compact.resolve())
    records = compact.get("records") or []
    if not records:
        raise RuntimeError("compact timeline has no records")
    people, scenes, props, default_scene = read_asset_rules(args.assets.resolve())
    translations, translations_by_dialogue_id, translation_warnings = build_dialogue_translations(compact, read_external_translations(args.translations_json.resolve() if args.translations_json else None))
    config = {
        "episode": clean(compact.get("episode") or "01"),
        "default_scene": default_scene,
        "groups": build_groups(records, scenes, default_scene),
        "people_rules": people,
        "scene_rules": scenes,
        "prop_rules": props,
        "translations": translations,
        "translations_by_dialogue_id": translations_by_dialogue_id,
        "dialogue_timing": {"hard_limit_wps": 4.5},
        "dialogue_segmentation": {"seam_tolerance_seconds": 0.15},
        "scene_resolution": {
            "minimum_record_score": 2,
            "minimum_support_records": 1,
            "minimum_duration_share": 0.0,
            "on_conflict": "warn",
        },
        "source": {
            "generator": "build_expand_config_from_compact.py",
            "compact": str(args.compact.resolve()),
            "assets": str(args.assets.resolve()),
            "warnings": translation_warnings,
        },
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "complete", "groups": len(config["groups"]), "people_rules": len(people), "scene_rules": len(scenes), "prop_rules": len(props), "output": str(args.output.resolve())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
