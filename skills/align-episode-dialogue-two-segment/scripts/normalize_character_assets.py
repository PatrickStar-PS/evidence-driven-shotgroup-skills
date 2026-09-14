#!/usr/bin/env python3
"""Normalize episode character assets for video-only dialogue discovery."""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Any


ALIASES = {
    "asset_type": ("资产类型", "类型", "asset type", "type"),
    "chinese_name": ("资产中文原名", "中文名", "人物名称", "资产名", "chinese name", "name"),
    "english_name": ("资产本地化英文名", "英国本地化名", "英文名", "localized english name", "english name"),
    "description": ("原始中文解析描述", "资产详细描述", "详细描述", "人物描述", "外观/身份描述", "description"),
    "episode": ("出现集号", "集数", "episode", "ep"),
}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def key(value: Any) -> str:
    return re.sub(r"[\s_\-/（）()：:]+", "", clean(value).lower())


def choose(headers: list[str], aliases: tuple[str, ...]) -> str | None:
    keyed = {key(item): item for item in headers}
    for alias in aliases:
        if key(alias) in keyed:
            return keyed[key(alias)]
    for alias in aliases:
        needle = key(alias)
        for normalized, original in keyed.items():
            if needle and (needle in normalized or normalized in needle):
                return original
    return None


def rows_from_file(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, list):
            return [dict(item) for item in data]
        for name in ("rows", "assets", "records"):
            if isinstance(data.get(name), list):
                return [dict(item) for item in data[name]]
        raise RuntimeError(f"No asset row list found in {path}")
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            return [dict(item) for item in csv.DictReader(stream, delimiter=delimiter)]
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required for XLSX input") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        result: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                headers = [clean(value) or f"column_{index + 1}" for index, value in enumerate(next(iterator))]
            except StopIteration:
                continue
            for values in iterator:
                record = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
                if any(clean(value) for value in record.values()):
                    result.append(record)
        return result
    raise RuntimeError(f"Unsupported asset input type: {path.suffix}")


def episode_matches(value: Any, wanted: str) -> bool:
    text = clean(value)
    if not text:
        return True
    numbers = {str(int(item)) for item in re.findall(r"\d+", text)}
    return str(int(wanted)) in numbers


def is_character(value: Any) -> bool:
    text = key(value)
    return not text or text in {"人物", "角色", "character", "person"}


def is_base_identity(chinese_name: str, english_name: str, description: str) -> bool:
    if "base identity" in description.lower():
        return True
    state_match = re.search(r"[-—]\s*([^-—]+)$", chinese_name)
    if state_match and state_match.group(1).strip() in {"默认态", "基础态", "本体"}:
        return True
    return not re.search(r"[-—]\s*[^-—]+$", chinese_name)


def base_chinese_name(value: str) -> str:
    value = re.sub(r"[（(]候选新增[）)]", "", clean(value)).strip()
    return re.split(r"[-—]", value, maxsplit=1)[0].strip()


def base_english_name(value: str) -> str:
    return re.split(r"\s*[-—]\s*", clean(value), maxsplit=1)[0].strip()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--assets-workbook", type=Path, required=True)
    parser.add_argument("--episode", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    episode = str(args.episode).zfill(2)
    source_path = args.assets_workbook.resolve()
    if source_path.suffix.lower() == ".json":
        source_payload = json.loads(source_path.read_text(encoding="utf-8-sig"))
        if isinstance(source_payload, dict) and isinstance(source_payload.get("characters"), list):
            source_episode = str(source_payload.get("episode") or episode).zfill(2)
            if source_episode != episode:
                raise RuntimeError(f"Normalized character JSON episode {source_episode} does not match requested episode {episode}")
            characters = [dict(item) for item in source_payload["characters"] if isinstance(item, dict)]
            if not characters or any(not clean(item.get("chinese_name")) for item in characters):
                raise RuntimeError("Normalized character JSON has no usable Chinese character identities")
            payload = dict(source_payload)
            payload["schema_version"] = "2.1-character-input-with-wardrobe-states"
            payload["episode"] = episode
            payload["upstream_asset_source"] = source_payload.get("asset_source")
            payload["asset_source"] = str(source_path)
            payload["characters"] = characters
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            print(json.dumps({"episode": episode, "characters": len(characters), "output": str(args.output.resolve()), "mode": "normalized_json_passthrough"}, ensure_ascii=False))
            return 0
    rows = rows_from_file(source_path)
    if not rows:
        raise RuntimeError("Asset input has no rows")
    headers = list(rows[0])
    mapping = {field: choose(headers, aliases) for field, aliases in ALIASES.items()}
    if not mapping["chinese_name"]:
        raise RuntimeError(f"Cannot find a character-name column; headers={headers}")

    matched_rows: list[dict[str, str]] = []
    for row in rows:
        if mapping["asset_type"] and not is_character(row.get(mapping["asset_type"])):
            continue
        if mapping["episode"] and not episode_matches(row.get(mapping["episode"]), episode):
            continue
        chinese = clean(row.get(mapping["chinese_name"]))
        english = clean(row.get(mapping["english_name"])) if mapping["english_name"] else ""
        description = clean(row.get(mapping["description"])) if mapping["description"] else ""
        if chinese:
            matched_rows.append({"chinese_name": chinese, "english_name": english, "description": description})

    characters: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for item in matched_rows:
        chinese, english, description = item["chinese_name"], item["english_name"], item["description"]
        if not chinese or not is_base_identity(chinese, english, description):
            continue
        chinese = base_chinese_name(chinese)
        english = base_english_name(english)
        if (chinese, english) in seen:
            continue
        seen.add((chinese, english))
        characters.append({
            "character_id": f"char-{len(characters) + 1:03d}",
            "chinese_name": chinese,
            "english_name": english,
            "description": description,
            "recognition_aliases": [chinese, english] if english else [chinese],
            "wardrobe_states": [],
        })

    by_chinese = {key(item["chinese_name"]): item for item in characters}
    by_english = {key(item["english_name"]): item for item in characters if item.get("english_name")}
    for item in matched_rows:
        chinese, english, description = item["chinese_name"], item["english_name"], item["description"]
        if is_base_identity(chinese, english, description):
            continue
        owner = by_chinese.get(key(base_chinese_name(chinese))) or by_english.get(key(base_english_name(english)))
        if owner is None:
            # Episode-scoped candidate rows are still valid recognition cards.
            candidate_chinese = base_chinese_name(chinese)
            candidate_english = base_english_name(english)
            marker = (candidate_chinese, candidate_english)
            if marker in seen:
                owner = next((candidate for candidate in characters if (candidate["chinese_name"], candidate["english_name"]) == marker), None)
            else:
                seen.add(marker)
                owner = {
                    "character_id": f"char-{len(characters) + 1:03d}",
                    "chinese_name": candidate_chinese,
                    "english_name": candidate_english,
                    "description": description,
                    "recognition_aliases": [value for value in (candidate_chinese, candidate_english, chinese, english) if value],
                    "wardrobe_states": [],
                    "candidate_asset": True,
                }
                characters.append(owner)
                by_chinese[key(candidate_chinese)] = owner
                if candidate_english:
                    by_english[key(candidate_english)] = owner
        state = {
            "chinese_asset_name": re.sub(r"[（(]候选新增[）)]", "", chinese).strip(),
            "english_asset_name": english,
            "description": description,
        }
        if state not in owner["wardrobe_states"]:
            owner["wardrobe_states"].append(state)
        for alias in (state["chinese_asset_name"], state["english_asset_name"]):
            if alias and alias not in owner["recognition_aliases"]:
                owner["recognition_aliases"].append(alias)
    if not characters:
        raise RuntimeError("No base character assets matched this episode")
    payload = {
        "schema_version": "2.1-character-input-with-wardrobe-states",
        "episode": episode,
        "asset_source": str(args.assets_workbook.resolve()),
        "characters": characters,
        "warnings": [],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"episode": episode, "characters": len(characters), "output": str(args.output.resolve())}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
