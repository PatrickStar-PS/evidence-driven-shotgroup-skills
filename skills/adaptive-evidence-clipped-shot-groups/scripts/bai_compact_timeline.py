#!/usr/bin/env python3
from __future__ import annotations

import argparse
import base64
import hashlib
import json
import random
import re
import time
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

from timeline_contract import build_boundary_contract, ensure_dialogue_contract


PROMPT_CONTRACT_VERSION = "3.9-mouth-motion-speech-sync"


def episode_matches(value: object, episode: str) -> bool:
    episode_number = int(str(episode))
    return re.search(rf"(?<!\d)EP0*{episode_number}(?!\d)", str(value or ""), flags=re.I) is not None


def episode_asset_reference_cards(path: Path, episode: str) -> list[str]:
    """Load every current-episode person, scene, and prop asset for relay matching."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    required = {"资产类型", "资产中文原名", "资产本地化英文名", "原始中文解析描述", "出现集号"}
    if not required.issubset(headers):
        raise RuntimeError(f"Asset workbook is missing required columns: {sorted(required - set(headers))}")
    positions = {name: headers.index(name) for name in required}
    cards: list[str] = []
    for row in rows:
        asset_type = str(row[positions["资产类型"]] or "").strip()
        if asset_type not in {"人物", "场景", "道具"}:
            continue
        if not episode_matches(row[positions["出现集号"]], episode):
            continue
        name = str(row[positions["资产中文原名"]] or "").strip()
        localized_name = str(row[positions["资产本地化英文名"]] or "").strip()
        description = re.sub(r"\s+", " ", str(row[positions["原始中文解析描述"]] or "").strip())
        if name and localized_name:
            cards.append(f"- [{asset_type}] {name}｜{localized_name}｜{description or '无原始中文解析描述'}")
    workbook.close()
    return cards


def image_data_url(path: Path) -> str:
    suffix = path.suffix.lower()
    mime = "image/png" if suffix == ".png" else "image/jpeg"
    return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode('ascii')}"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_evidence_pack(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    manifest_path = path.resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8-sig"))
    pages = manifest.get("relay_pages") or []
    if not isinstance(pages, list) or not pages:
        raise RuntimeError("evidence pack has no relay_pages")
    allowed_roles = {"range_boundaries", "action_events", "high_motion_triplets", "shot_space", "person_trajectory", "multi_person_geometry"}
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in pages:
        role = str(item.get("role") or "")
        if role not in allowed_roles:
            raise RuntimeError(f"unsupported evidence role: {role!r}")
        if role in seen:
            raise RuntimeError(f"duplicate evidence role: {role}")
        source = Path(str(item.get("file") or ""))
        source = source if source.is_absolute() else manifest_path.parent / source
        source = source.resolve()
        if not source.exists():
            raise RuntimeError(f"evidence page does not exist: {source}")
        expected = str(item.get("sha256") or "")
        actual = file_sha256(source)
        if expected and expected != actual:
            raise RuntimeError(f"evidence checksum mismatch for {role}: {source}")
        result.append({"role": role, "path": source, "sha256": actual, "summary": item.get("summary") or {}})
        seen.add(role)
    required = {"range_boundaries", "action_events", "high_motion_triplets", "shot_space", "person_trajectory"}
    missing = sorted(required - seen)
    if missing:
        raise RuntimeError(f"evidence pack is missing required roles: {missing}")
    gate = manifest.get("pose_gate") or {}
    if bool(gate.get("pose_enabled")) != ("multi_person_geometry" in seen):
        raise RuntimeError("pose gate and multi_person_geometry attachment disagree")
    return result, manifest


def atomic_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def read_keys(path: Path) -> list[str]:
    result: list[str] = []
    for raw in path.read_text(encoding="utf-8-sig").replace(",", "\n").splitlines():
        value = raw.strip()
        if not value or value.startswith("#"):
            continue
        if "=" in value:
            value = value.split("=", 1)[1].strip().strip('"').strip("'")
        if value and value not in result:
            result.append(value)
    if not result:
        raise RuntimeError("中转站 key pool is empty")
    return result


def response_text(data: dict[str, Any]) -> str:
    choices = data.get("choices")
    message = choices[0].get("message") if isinstance(choices, list) and choices else None
    content = message.get("content") if isinstance(message, dict) else None
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        return "".join(part.get("text", "") for part in content if isinstance(part, dict)).strip()
    raise RuntimeError("中转站 response content is empty")


def extract_json(text: str) -> dict[str, Any]:
    stripped = text.strip()
    if stripped.startswith("```"):
        stripped = re.sub(r"^```(?:json)?\s*", "", stripped, flags=re.I)
        stripped = re.sub(r"\s*```$", "", stripped)
    try:
        value = json.loads(stripped)
    except json.JSONDecodeError:
        start, end = stripped.find("{"), stripped.rfind("}")
        if start < 0 or end <= start:
            raise RuntimeError("中转站 response does not contain a JSON object")
        value = json.loads(stripped[start : end + 1])
    if not isinstance(value, dict):
        raise RuntimeError("中转站 response JSON root must be an object")
    return value


def load_clip_mapping(path: Path, range_id: str) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    matches = [item for item in data.get("ranges") or [] if str(item.get("id")) == range_id]
    if len(matches) != 1:
        raise RuntimeError(f"clip mapping must contain exactly one range named {range_id!r}")
    item = matches[0]
    required = {
        "clip_start_seconds", "clip_end_seconds", "logical_start_seconds", "logical_end_seconds", "absolute_offset_seconds",
        "core_start_clip_seconds", "core_end_clip_seconds",
    }
    missing = sorted(required - set(item))
    if missing:
        raise RuntimeError(f"clip mapping {range_id} is missing fields: {missing}")
    return item


def clip_transport_note(mapping: dict[str, Any] | None) -> str:
    if not mapping:
        return ""
    return (
        "TRANSPORT_MODE=overlap_clipped_video. The attached video is not the whole episode. "
        f"Its local 0.000s equals original absolute {float(mapping['clip_start_seconds']):.3f}s; "
        f"the attached clip covers original {float(mapping['clip_start_seconds']):.3f}-"
        f"{float(mapping['clip_end_seconds']):.3f}s. Analyze and output only the logical core original "
        f"{float(mapping['logical_start_seconds']):.3f}-{float(mapping['logical_end_seconds']):.3f}s, "
        f"which appears at clip-local {float(mapping['core_start_clip_seconds']):.3f}-"
        f"{float(mapping['core_end_clip_seconds']):.3f}s. Pre/post-roll is context only. Prefer original "
        "absolute seconds in every returned timestamp. If this range has seam_owner=start, audit the "
        "logical start boundary from the visible pre-roll and core frames; never assume a file seam is a cut."
    )


def normalize_clipped_timestamps(
    data: dict[str, Any], mapping: dict[str, Any] | None, tolerance: float = 0.20
) -> str:
    if not mapping:
        return "full_episode_absolute"
    records = sorted(data.get("records") or [], key=lambda item: float(item.get("start_seconds", 0.0)))
    if not records:
        raise RuntimeError("clipped response contains no records")
    absolute_start = float(mapping["logical_start_seconds"])
    absolute_end = float(mapping["logical_end_seconds"])
    local_start = float(mapping["core_start_clip_seconds"])
    local_end = float(mapping["core_end_clip_seconds"])
    first = float(records[0].get("start_seconds", 0.0))
    last = float(records[-1].get("end_seconds", first))
    absolute_error = abs(first - absolute_start) + abs(last - absolute_end)
    local_error = abs(first - local_start) + abs(last - local_end)
    if absolute_error <= tolerance * 2:
        mode = "model_absolute"
        offset = 0.0
    elif local_error <= tolerance * 2:
        mode = "local_shifted"
        offset = float(mapping["absolute_offset_seconds"])
    else:
        raise RuntimeError(
            "clipped response coverage matches neither absolute nor clip-local core: "
            f"first={first:.3f}, last={last:.3f}, absolute={absolute_start:.3f}-{absolute_end:.3f}, "
            f"local={local_start:.3f}-{local_end:.3f}"
        )
    if offset:
        for record in records:
            for key in ("start_seconds", "end_seconds"):
                if record.get(key) is not None:
                    record[key] = round(float(record[key]) + offset, 3)
            timestamps = record.get("anchor_evidence_timestamps")
            if isinstance(timestamps, list):
                record["anchor_evidence_timestamps"] = [round(float(value) + offset, 3) for value in timestamps]
        for item in data.get("boundary_audit") or []:
            for key in ("candidate_seconds", "observed_seconds"):
                if item.get(key) is not None:
                    item[key] = round(float(item[key]) + offset, 3)
    normalized_first = float(records[0]["start_seconds"])
    normalized_last = float(records[-1]["end_seconds"])
    if abs(normalized_first - absolute_start) > tolerance or abs(normalized_last - absolute_end) > tolerance:
        raise RuntimeError("clipped response does not cover the complete logical core after timestamp normalization")
    return mode


def load_dialogue_injection(
    path: Path,
    ledger_path: Path,
    episode: str,
    range_start: float,
    range_end: float,
) -> dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if data.get("schema_version") not in {"2.0-dialogue-injection", "2.1-dialogue-injection-delivery"}:
        raise RuntimeError("dialogue injection must use schema_version 2.0-dialogue-injection or 2.1-dialogue-injection-delivery")
    if str(data.get("episode")).zfill(2) != str(episode).zfill(2):
        raise RuntimeError("dialogue injection episode does not match --episode")
    source = data.get("source_ledger")
    if not isinstance(source, dict) or not str(source.get("sha256") or "").strip():
        raise RuntimeError(
            "dialogue injection is not source-bound; rebuild it from the current validated dialogue ledger"
        )
    ledger_path = ledger_path.resolve()
    current_hash = file_sha256(ledger_path)
    injected_hash = str(source.get("sha256") or "").strip().lower()
    if current_hash.lower() != injected_hash:
        raise RuntimeError(
            "stale dialogue injection: source ledger fingerprint differs from --dialogue-ledger; rebuild before relay analysis"
        )
    ledger = json.loads(ledger_path.read_text(encoding="utf-8-sig"))
    if str(ledger.get("episode")).zfill(2) != str(episode).zfill(2):
        raise RuntimeError("dialogue ledger episode does not match --episode")
    if str(source.get("schema_version") or "") != str(ledger.get("schema_version") or ""):
        raise RuntimeError("stale dialogue injection: source ledger schema changed; rebuild before relay analysis")
    context_events: list[dict[str, Any]] = []
    analysis_events: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in data.get("events") or []:
        event_id = str(item.get("dialogue_id") or "").strip()
        if not event_id or event_id in seen:
            raise RuntimeError(f"duplicate or empty injected dialogue_id: {event_id!r}")
        seen.add(event_id)
        start = float(item["start_seconds"])
        end = float(item["end_seconds"])
        if end <= start:
            raise RuntimeError(f"invalid injected dialogue range for {event_id}")
        event = {
            "dialogue_id": event_id,
            "start_seconds": round(start, 3),
            "end_seconds": round(end, 3),
            "speaker": str(item.get("speaker_chinese_name") or "unclear"),
            "text": str(item.get("chinese_text") or ""),
            "basis": "dialogue_injection",
            "speaker_evidence": "locked_dialogue_ledger",
        }
        if isinstance(item.get("delivery"), dict):
            event["delivery"] = json.loads(json.dumps(item["delivery"], ensure_ascii=False))
        context_events.append(event)
        if start < range_end and end > range_start:
            analysis_events.append(event)
    return {
        "schema_version": "2.4-dialogue-injection-delivery-locked" if any(item.get("delivery") for item in context_events) else "2.3-dialogue-injection-locked",
        "analysis_range": {"start_seconds": range_start, "end_seconds": range_end},
        "analysis_events": analysis_events,
        "context_events": context_events,
        "source_ledger": {
            "name": ledger_path.name,
            "sha256": current_hash,
            "schema_version": ledger.get("schema_version"),
        },
    }


def hydrate_dialogue_injection(data: dict[str, Any], injection: dict[str, Any]) -> None:
    locked = injection.get("analysis_events") or []
    locked_by_id = {item["dialogue_id"]: item for item in locked}
    for record in data.get("records") or []:
        record_start = float(record.get("start_seconds", 0.0))
        record_end = float(record.get("end_seconds", record_start))
        visibility_by_id: dict[str, str] = {}
        for placement in record.get("dialogue_placements") or []:
            if not isinstance(placement, dict):
                continue
            event_id = str(placement.get("dialogue_id") or "")
            visibility = str(placement.get("speaker_visibility") or "unclear")
            if event_id in locked_by_id and visibility in {"on_screen", "off_screen", "unclear"}:
                visibility_by_id[event_id] = visibility
        for utterance in record.get("dialogue_utterances") or []:
            if not isinstance(utterance, dict):
                continue
            event_id = str(utterance.get("dialogue_id") or "")
            visibility = str(utterance.get("speaker_visibility") or "unclear")
            if event_id in locked_by_id and visibility in {"on_screen", "off_screen", "unclear"}:
                visibility_by_id.setdefault(event_id, visibility)
        overlaps = [item for item in locked if float(item["start_seconds"]) < record_end and float(item["end_seconds"]) > record_start]
        record["dialogue_event_ids"] = [item["dialogue_id"] for item in overlaps]
        record["dialogue_utterances"] = [
            {
                "dialogue_id": item["dialogue_id"],
                "start_seconds": round(max(float(item["start_seconds"]), record_start), 3),
                "end_seconds": round(min(float(item["end_seconds"]), record_end), 3),
                "speaker": item["speaker"],
                "speaker_visibility": visibility_by_id.get(item["dialogue_id"], "unclear"),
                "text": item["text"],
                "basis": "dialogue_injection",
                "delivery": json.loads(json.dumps(item.get("delivery") or {}, ensure_ascii=False)),
            }
            for item in overlaps
        ]
        record["speaker"] = "、".join(dict.fromkeys(item["speaker"] for item in overlaps))
        record["chinese_dialogue"] = "；".join(item["text"] for item in overlaps)
    data["dialogue_events"] = locked
    data["dialogue_injection_contract"] = {
        "version": injection.get("schema_version") or "2.3-dialogue-injection-locked",
        "event_count": len(locked),
        "semantic_fields_locked": ["dialogue_id", "start_seconds", "end_seconds", "speaker", "text", "delivery"],
        "source_ledger": injection.get("source_ledger") or {},
    }


def prompt(
    episode: str,
    duration: float,
    boundary_contract: dict[str, Any] | None = None,
    range_start: float = 0.0,
    range_end: float | None = None,
    has_context_sheet: bool = False,
    has_pose_evidence: bool = False,
    evidence_roles: list[str] | None = None,
    asset_cards: list[str] | None = None,
    dialogue_injection: dict[str, Any] | None = None,
    continuity_injection: dict[str, Any] | None = None,
) -> str:
    range_end = duration if range_end is None else range_end
    boundary_note = ""
    if boundary_contract:
        candidate_text = ", ".join(
            f"{float(item['seconds']):.3f}({item['tier']},score={item.get('score')})"
            for item in boundary_contract.get("candidates") or []
        )
        interval_text = "\n".join(
            f"- {item['evidence_interval_id']}=[{float(item['start_seconds']):.3f},{float(item['end_seconds']):.3f}) "
            f"anchors={float(item['start_anchor_seconds']):.3f}/{float(item['middle_anchor_seconds']):.3f}/{float(item['end_anchor_seconds']):.3f}"
            for item in boundary_contract.get("evidence_intervals") or []
        )
        locked_text = ", ".join(
            f"{item['interval_id']}=[{float(item['start_seconds']):.3f},{float(item['end_seconds']):.3f})"
            for item in boundary_contract.get("locked_intervals") or []
        )
        boundary_note = f"""
本地逐帧差异得到以下候选切点：[{candidate_text}]。tier=locked表示高分时间锚点：若裁决为keep_cut或keep_reframe，必须保留原时间，不能前后移动；若首中尾证据明确证明同一连续镜头，仍可merge_false。tier=advisory同样须裁决，但不是事实真值。候选之外若存在真实切镜仍须补充。
所有候选证据区间及首/中/尾锚点如下：
{interval_text}
高分时间锚点形成的父区间：{locked_text}
每条record必须填写所属父interval_id及覆盖的evidence_interval_ids。若任一候选被merge_false，一条record可覆盖相邻多个C区间；若保留则必须在对应C区间边界拆开。逐个查看当前record覆盖证据区间的首/中/尾锚点并填写start_anchor_subject、middle_anchor_subject、end_anchor_subject和anchor_evidence_timestamps。primary_visible_subject必须描述当前[start_seconds,end_seconds)内部画面，严禁把end_seconds之后下一镜的人物提前填入当前区间。每次boundary_audit仍须填写切点前后主体与same_shot_evidence；台词、字幕连续不能证明视觉连续。
"""
    role_text = ",".join(evidence_roles or [])
    context_note = "" if not has_context_sheet else f"""
同时提供本解析范围的角色标注证据页：[{role_text}]。PAGE_ROLE=range_boundaries只用于范围入口和出口；PAGE_ROLE=action_events按每个候选镜头提供动作开始、变化/峰值和结束的未标注原始帧；PAGE_ROLE=high_motion_triplets只强调最强运动镜头的前—中—后变化；PAGE_ROLE=shot_space用放大原始像素检查画面站位、鼻尖、耳朵、肩线、裁切、占比、遮挡、物理前后与光学清晰度。先按时间和interval_id归档，不得把下一镜帧解释成当前镜头。逐个比较每名人物在首、中、尾锚点是否可见，并填写anchor_presence；若人物只在中间出现或在首尾间进入/离开，必须回看视频判断其进入、经过、交错、短暂遮挡或离开过程，不能把该人物降格为静态边缘人物或笼统写“动作不明”。对每个人的character_action和镜头visible_action，尽可能写成“起始状态→可见过程/峰值→结束状态”的动作链，不用一句概括吞掉过程。模型仍必须以裁剪视频的原始画面与声音作为最高证据；候选帧不是切镜真值，不得把说话人当成可见人物。
"""
    trajectory_note = "" if "person_trajectory" not in (evidence_roles or []) else """
另附PAGE_ROLE=person_trajectory人物轨迹总页。检测框、脚点路径、Sxxx-Txx编号、进入/离开、短时缺失和交错事件均为单个真实镜头内部的几何证据；编号在每次cut处重置，绝不能跨镜头绑定，也不能直接对应人物姓名。轨迹页只用于发现人物进入、经过、交错、停止、接近、离开或短时遮挡的候选过程，并辅助核对anchor_presence与三阶段character_action。它不能判断身份、动作意图、人物朝向、物理景深、接触关系或台词说话人。相机补偿轨迹是估计值；若与裁剪视频或未标注原帧冲突，以视频和原帧为准。
"""
    pose_note = "" if not has_pose_evidence else """
另附一张PAGE_ROLE=multi_person_geometry的JointBDOE＋MediaPipe Pose Lite骨架姿态页。绿色框/骨架表示Pose Lite验证到足够关键点，橙色框表示未验证或证据不足；P编号仅在单帧内有效，不能跨帧绑定人物。该页只允许辅助发现多人、standing/sitting等粗姿态、身体关节几何、肢体弯曲和姿态变化。JointBDOE body angle与骨架肩线都不能单独决定躯干朝画面左/右；torso_orientation必须回看未标注原片中的胸口、衣襟、肩部透视和身体表面。该页不能判断人物身份、头部朝向、视线目标、面对关系、物理景深、动作意图或台词。若骨架与原视频像素、未标注联系图冲突，必须以原视频和未标注帧为准；裁切、遮挡或虚焦导致骨架缺失时，不得据此删除可见人物，证据不足写unclear。
"""
    asset_note = ""
    if asset_cards:
        asset_note = """
以下是按资产表“出现集号”筛出的本集完整中英双语资产参考卡，包含人物、场景和道具，格式为“[资产类型] 中文源资产名｜英国本地化规范名｜中文识别描述”。参考卡只用于匹配视频中已经直接看到的内容，不证明资产必然出现，也不得压过视频证据。
先观察后匹配。人物必须先独立清点前景、中景、背景与边缘人体，再按服装、发型、年龄、性别和饰品比对；明显吻合时identity与asset_hints使用精确中文源资产名，appearance写asset_bound。场景依据scene_observation与fixed_scene_evidence比对；道具依据primary_visible_subject、visual_details、transient_visual_details和prop_continuity逐件比对。人物、场景或道具只要明显命中参考卡，都必须把精确中文源资产名写入asset_hints；道具的prop_continuity.identity也优先使用精确中文源资产名。英国名仅作固定映射参考，不写入中文compact identity。部分吻合时保留观察名称并降低confidence；不吻合时不得强行绑定。场景和物体永远不得进入visible_characters。
本集完整资产参考卡：
""" + "\n".join(asset_cards) + "\n"
    orientation_note = """
人物朝向属于最高优先级视觉事实。逐人按以下顺序判断并在body_orientation或relative_blocking.evidence中保留简短证据：耳朵/后脑到鼻尖的突出方向决定头部朝画面左或右；双肩线、胸口与衣襟轴线决定躯干；瞳孔单独决定视线。鼻尖向红色SCREEN LEFT边框突出时必须写头部朝画面左，绝不能因为该人物站在另一人物左侧就反推为朝右。鼻尖或肩线看不清时写unclear，不按剧情关系猜测。多人镜头必须反问他们是真的面对面，还是共同看向同一画外方向；只有双方头部或躯干轴线相向才写face_to_face。物理前后仅根据占画比例、贴边裁切、透视放大和遮挡判断；大面积贴边裁切的虚焦人物优先视为极近前景，不得因虚焦写成背景。
"""
    orientation_note += """
水平机位只根据摄影机光轴相对场景透视轴、人物主运动轴和稳定构图的关系判断，不根据人物自身转头、低头、侧身或三分之四躯干朝向判断。人物沿画面纵深迎着摄影机走来或背向摄影机远去，且走廊/地面/墙线消失点位于其身后或画面中央时，即使人物因交谈而呈三分之四身体，仍写“2 正视”。只有摄影机光轴明显位于人物主运动轴侧方、主要观察人物横向经过画面并持续呈现侧面轮廓时，才写“3 侧视”；单帧侧脸或“前方偏左/偏右”不足以证明侧视。硬切、景别变化、单人变双人和背景改变只决定boundary/transition，不能自动改变camera_view；硬切前后可以同为“2 正视”。camera_view_evidence必须分别说明摄影机光轴、人物运动方向和场景透视消失点，禁止拿torso_orientation或head_orientation充当机位依据。“4 反打”只用于越过对话轴线后从对方一侧回看的反向机位。
对anchor_presence为start=false、middle=true、end=false的人物，character_action必须严格写成三个由“→”连接的可见阶段：“从哪侧/哪层进入→相对哪个主体如何经过、交错或造成遮挡→从哪侧/向哪层离开或在尾锚点消失”。不得只写“路人走过”“快速走过”或“边缘可见”。如果某一阶段确实不可见，明确写“进入前不在画面/离开过程不可见/尾锚点已消失”，不能省略阶段。
"""
    continuity_note = """
每个cut或reframe先单独填写screen_order_left_to_right：按证据页底部红色SCREEN LEFT到蓝色SCREEN RIGHT的观众坐标，分别记录start/middle/end锚点中可见人物从左到右的顺序。随后才填写每个人的screen_position、relative_blocking.horizontal_relation和spatial_positions；这三处必须从同一个screen_order_left_to_right派生，不得各自猜测。身份或重叠不清时写unclear，禁止左右反转。每条record还必须填写position_transition={type,description,evidence_timestamps}。type只能是stable、crosses_behind、crosses_in_front、reblocked_after_cut或unclear。同一连续镜头内start/middle/end共同人物的左右顺序发生反转时，必须写清谁从谁的前方或后方经过以及最终到达哪一侧；跨硬切发生的重新站位写在硬切后的record中，type=reblocked_after_cut，绝不能虚构为连续横穿动作。
如果同一record的start/middle/end顺序、人物所处画面侧或前后层发生变化，必须把该变化写成可见运动：在对应人物character_action和镜头visible_action中写清“从画面哪侧/哪层→经过谁的前方或后方→到画面哪侧/哪层”。screen_position不得把整段压成一个固定站位；只有三个锚点位置确实不变时才可写静态位置。人物暂时被主体遮挡后换侧重现，按连续经过/交错/遮挡处理，不能当成瞬移或新人物。
位置字段必须做反向一致性检查：只要任一visible_characters.character_action或visible_action写出“左侧/左后方→右侧/右后方”或反向换侧，position_transition.type绝不能写stable；必须依据人物从主体前方还是后方经过，填写crosses_in_front或crosses_behind，并让screen_order_left_to_right的首/中/尾锚点表现出该人物在可见时的真实换侧。若换侧中间被主体完全遮挡，middle锚点可暂不列该人物，但description必须明确“被谁遮挡后在另一侧重现”。反之，若position_transition.type=stable，所有人物动作不得出现跨越主体换侧的描述。
screen_order_left_to_right与anchor_presence必须逐锚点一一相符：anchor_presence.start=false的人物不得写入start，middle=false不得写入middle，end=false不得写入end；反向也一样，锚点列表中的每个人在对应anchor_presence必须为true。比较首中尾时，以脸、眼镜、发型以及持续持有的文件夹、包或其他道具等稳定外观线索追踪同一人物，不能因其短暂被主体遮挡、换到另一侧或轨迹编号改变而当成不同人物。尤其要逐个核对“中锚点位于主体一侧、尾锚点携带同一物件出现在另一侧”的人物，这通常是连续绕行换位而不是stable。
临时道具按物理实例追踪。每个持续存在的物件使用稳定instance_id，并填写visible_instance_count、appearance_invariants与exclusive_holder_after。翻面、展开、合拢或露出不同颜色的封面/内衬仍是同一实例；只有同一帧明确同时可见两件彼此分离的物品时，才允许visible_instance_count大于1或创建第二个instance_id。交接动作必须写成单一实例的三阶段：交接前仅交出者持有→交接中双方短暂接触同一物品→交接后仅接收者持有，交出者对应手部为空。交接完成后禁止让双方各自保留一件。物理裁剪片的pre-roll属于核心起点的前镜连续性证据；pre-roll已出现的道具不得在logical core起点写“硬切后首次揭示”或continuity_from_previous=none。
"""
    if continuity_injection:
        continuity_note += """
下面提供“上一逻辑核心尾状态”，它来自相邻范围上一条模型结果的结构化尾帧，仅用于接缝连续性，不是对当前画面的语义改写。当前范围第一条record必须先在pre-roll和原视频中核对它：若未看到明确交接、放下、消失或第二件独立物品，则沿用相同instance_id、同一物品身份、上一持有者和开合状态，source写上一核心继承，continuity_from_previous明确引用上一范围；不得因裁剪边界、硬切、翻面或内外颜色不同把它改成首次出现或复制第二件。若当前画面明确推翻上一状态，按当前视频写事实并在warnings说明直接证据。
上一逻辑核心尾状态：
""" + json.dumps(continuity_injection, ensure_ascii=False, separators=(",", ":")) + "\n"
    dialogue_requirement = """8. 先在根对象dialogue_events中建立独立台词事件，每句只建一次，填写dialogue_id、真实起止时间、speaker、text、basis和speaker_evidence=audio_voice|lip_sync|context|unclear。然后各视觉record通过dialogue_event_ids引用事件，并在dialogue_utterances中保留兼容项；兼容项必须携带同一个dialogue_id及本镜头speaker_visibility。跨切镜持续的同一句台词必须保持同一dialogue_id和speaker，反应镜头只能把speaker_visibility改为off_screen，不能重新判断说话人或重复创建事件。只有明确听到第二个人复述，才允许相同文本使用不同speaker，并填写speaker_change_evidence。chinese_dialogue按本record实际覆盖的事件文本顺序连接。若听不清但烧录字幕清楚，可照录并警告。画面文字须区分人物标签、剧情字幕、文件、地图、车牌和英文烧录字幕。"""
    dialogue_qa = "凡mouth_dynamics显示人物开口说话，或on_screen_text包含台词字幕，该句都必须出现在dialogue_utterances和chinese_dialogue中。"
    dialogue_note = ""
    if dialogue_injection:
        dialogue_note = """
整集台词已由独立台词解析锁定。下面analysis_events是本范围唯一允许使用的台词真值；context_events只用于范围边缘上下文。严禁重新听写、纠正文案、修改说话人、改动绝对时间、改写delivery、创建新dialogue_id或因为切镜拆分台词。delivery中的语速、语调、音色、情绪、节奏与停顿均来自整集音频，本阶段不得依据画面重新判断。剧情、音频、字幕和口型在本阶段只用于确定每个视觉record中说话人是on_screen、off_screen或unclear，以及理解切镜和人物反应。
每条record只输出dialogue_placements=[{"dialogue_id":"EP09-D0001","speaker_visibility":"on_screen|off_screen|unclear"}]，只列与record时间相交的注入事件；dialogue_utterances输出空数组，speaker与chinese_dialogue输出空字符串，根dialogue_events输出空数组，本地程序会按锁定台账确定性补齐。若看到疑似台词但注入台账没有对应事件，只在warnings写possible_unledgered_dialogue，不得自行增加台词。
源片烧录中文字幕只是锁定台词的视觉重复，不再OCR或逐字抄录；看到这类字幕时，on_screen_text固定写“台词字幕（与锁定台词同步）”，不要添加冒号或字幕正文。人物标签、文件、地图、车牌等非台词文字仍按画面事实记录。
character_action、visible_action与mouth_dynamics只描述可见身体动作、接触、嘴部开合和反应，不得引用、复述或改写台词文字；台词内容只由锁定台账补齐。
反应镜头必须执行“声音人物与画面人物解耦”：若本record相交的锁定台词说话人标为off_screen，且画面中没有属于可见人物的另一条on_screen锁定台词，则所有可见人物的mouth_state不得写speaking；character_action、visible_action、beat_summary、primary_visible_subject与mouth_dynamics不得出现“说话、讲话、开口、配合台词、念出台词”等发言含义，只能写实际可见的表情、视线、呼吸和肢体反应。烧录字幕覆盖在某人身上、单帧嘴部张开、笑容露齿或手势动作，都不能证明该人正在说字幕台词；必须以锁定说话人与speaker_visibility为准。只有台账中存在该可见人物自己的on_screen事件时，才允许写speaking及发言动作。
逐record先从dialogue_placements建立on_screen_locked_speakers集合，再填写嘴部和动作字段：可见人物identity不在该集合中时，不得写mouth_state=speaking，也不得使用任何发言动词；集合为空时，本record所有可见人物只能写not_speaking或unclear，并把嘴部变化描述为“嘴部短暂张开/闭合但无锁定发言依据”，不能解释成台词。输出前反向扫描五处字段：beat_summary、primary_visible_subject、visible_characters[].character_action、visible_action、mouth_dynamics；其中任何发言含义都必须能对应同一record内该人物自己的on_screen锁定事件，否则删除发言含义并改写为可见反应。
嘴部可见动作与发言归属必须分字段记录。每个可见人物都填写mouth_visual_action（只写闭合、短暂张开后闭合、连续开合、微笑露齿、抿嘴或unclear等直接可见变化）、speech_sync_status=matched_on_screen_dialogue|off_screen_audio_reaction|non_speech_mouth_motion|no_visible_mouth_motion|unclear，以及mouth_action_evidence（至少一个时间点或首/中/尾对比证据）。嘴部张开、露齿、笑容或单帧口型只能进入mouth_visual_action，不能单独推出speaking。mouth_state=speaking只允许同时满足：该人物拥有本record内自己的on_screen锁定台词，并且连续帧口型呈发声式开合；此时speech_sync_status必须为matched_on_screen_dialogue。若声音来自画外人物而画面人物嘴部有变化，画面人物仍写mouth_state=not_speaking，speech_sync_status=off_screen_audio_reaction，并把变化写成可见反应动作。
dialogue_placements的speaker_visibility必须通过“锁定说话人身份—可见人物身份”一致性检查：只有锁定事件的speaker_label或其明确基础身份实际出现在本record的visible_characters中，才允许标为on_screen；不存在就必须标为off_screen。绝对禁止把A人物（例如助理）的锁定台词，因为B人物（例如胖子）恰好张嘴或做手势，就转贴给B人物或把A虚构成画内。输出前逐条核对：每个on_screen placement都能在visible_characters找到同一身份；每个matched_on_screen_dialogue也能找到同一人物自己的on_screen placement，否则一律改为off_screen_audio_reaction或unclear，并删除发言语义。
手部动作必须按接触形态描述：只有双掌掌面持续贴合才可写“双手合十”；手指相扣、指尖相碰、搓手、捏指或两手在胸前比划必须使用对应中性动作，不得升级为祈求式合十。
锁定台词注入：
""" + json.dumps(dialogue_injection, ensure_ascii=False, separators=(",", ":")) + "\n"
        dialogue_requirement = """8. 台词语义与delivery声音表现均已由注入台账锁定。不得重新发现、转写、翻译、纠正、拆分、合并、改派台词或改写声音表现；只为每条视觉record填写已注入dialogue_id及speaker_visibility。视觉切镜仍根据画面、动作、构图和剧情节奏判断，台词连续既不能证明视觉连续，也不能阻止真实切镜。"""
        dialogue_qa = "逐条检查dialogue_placements只引用注入ID，并覆盖所有与本record时间相交的锁定事件；不要重复输出台词正文。若事件为off_screen，复核画面反应人物没有被写成speaking、说话、讲话或开口；字幕和单帧张嘴不算说话证据。"
    rendered = f"""你是专业短剧拉片分析师。完整观看第{episode}集视频，总时长以本地检测值 {duration:.3f} 秒为准。本次只解析绝对时间范围 {range_start:.3f}—{range_end:.3f} 秒，所有时间必须保留相对原视频的绝对秒数，不得从0重新计时。范围外内容只用于理解上下文，不得输出记录。{context_note}{trajectory_note}{pose_note}{orientation_note}{continuity_note}{boundary_note}{asset_note}
{dialogue_note}
只做第一阶段的高密度事实时间轴，不做英国本地化，不写生成提示词；目标是让没有看过视频的人，仅凭记录即可准确复原每个镜头的画面与动作。

要求：
1. 从{range_start:.3f}秒连续覆盖到{range_end:.3f}秒，第一条记录必须从{range_start:.3f}开始，最后一条必须在{range_end:.3f}结束。每次硬切、正反打、侧向机位切换、独立反应、物件/手部插入特写、景别跳变、轴线或人物位置重置、明显重构图，以及连续运镜结束后形成新稳定构图时必须新建记录；不得出现大于0.10秒的空档、重叠或合并掉的短插镜。不要为了压缩输出而吞并镜头。
2. 每条只写直接可见或可听事实。scene_observation与fixed_scene_evidence可记录固定场景陈设用于第一阶段匹配，人物识别时也可使用服装特征；但一旦人物匹配到“人物-服装态”，appearance必须写asset_bound，其他动作与空间字段不得复述其服装颜色、材质、款式或配饰。资产匹配只决定人物名称和是否省略服装描述，绝不能决定人物是否出镜：未匹配资产的真实人物仍必须保留，按同一镜头内稳定顺序命名为“背景路人A/B/C”。visible_characters只能填写人，严禁把场景、建筑、车辆、喷泉、家具、灯具或其他物体写成人物。人物屏幕位置、姿态、朝向、支点、手部动作、人与物接触的动作链、首中尾人物出现状态、首末帧变化、表情视线、嘴部、声音和画面文字必须具体。initial_frame与final_frame只写人物/物体的位置、动作、表情、视线及瞬时状态，不写人物服装、固定装修或常设陈设。逐帧核对原视频、事件关键帧、运动三联帧和人物轨迹页；任何真实人体即使不说话、虚焦、占画tiny、只露局部、位于背景或画框边缘、只出现于中间锚点、只从主体身边经过，都必须进入visible_characters并写清可见运动过程。优先写“抓住领带→抽出→举到胸前查看”或“画面右侧进入→从主体身边交错经过→向画面后方离开”这类可复原动作链，避免“整理东西”“两人互动”“路人走过”等过度概括。
3. visual_details与fixed_scene_evidence仅作为第一阶段场景/资产识别证据，不直接进入最终生成提示词。另填transient_visual_details，只记录本镜头临时出现、且未被人物服装态或场景资产覆盖的剧情道具、物件状态和位置；不得复述固定装修、家具、墙面颜色、材质、常设陈设，也不得复述已匹配人物的衣服和配饰。凡被人物持有、携带、递交、接取、打开、放下、展示、拿出或使用的临时道具，还必须逐件写入prop_continuity：identity使用稳定名称；holder写当前持有者或none；support_contact写左/右手、双手、手臂夹持、桌面承托等直接接触；start_state、middle_action、end_state分别写首帧状态、中间物理动作和尾帧状态；source写道具从哪里进入本镜头；continuity_from_previous和continuity_to_next写与相邻record的继承关系。只要visible_characters.character_action、visible_action、contact_actions、initial_frame或final_frame中出现“手持/抱着/夹着/拿着/递出/接过/翻开/展示/使用”某物，该record的prop_continuity中就必须有同一道具，不能因为transient_visual_details写了“无”而省略。道具首次参与动作时必须交代来源；若上一镜已存在，下一镜必须继承同一道具的持有者、尺寸、开合状态与位置。硬切只改变镜头，不会清除上一镜已经可见的道具；只有该道具在硬切后的record内确实是全片首次可见时，才允许写“硬切后首次揭示”。除非画面明确显示从口袋、包、抽屉、画外递入或硬切后首次揭示，否则不得让道具凭空出现或消失。无法看到来源时写unclear并加入warnings，禁止脑补。generation_facts是唯一允许进入最终构图描述的结构：subject_lighting只写主体受光是否正常，异常时可写逆光、欠曝、过曝或被动态光源短时照亮；contrast_exposure只写整体明暗程度、曝光和反差；depth_of_field只能写shallow|deep|rack_focus|unclear；focus_transition只写焦点是否及如何转移；composition_change只写构图重心、前景遮挡或景别变化；dynamic_environment只写门窗开合、车辆运动、烟雾/雨水，以及灯光开启、熄灭、闪烁警灯、车灯扫过、屏幕光照到脸上等剧情性变化。固定灯带、吊灯、顶光、霓虹灯和彩色氛围灯由场景资产负责；“蓝色、红色、霓虹、灯带持续发光”等静态环境光不得进入generation_facts。generation_facts严禁出现人物服装、配饰、墙面、地面、家具、灯具造型、装修材质或常设陈设。不得因为临时物件不是剧情主角而省略。
4. transition精确说明硬切、反打、横甩/纵甩、推拉模糊、叠化、淡入淡出、闪白闪黑、变焦、镜面反射、遮挡/擦镜、人物出入画或无明显转场。凡转场前稳定画面、转场运动/模糊区、转场后稳定画面的主体或构图不同，应分别建立记录并写清开始和结束；同一稳定构图内的字幕变化、眨眼、转头、抬手、普通走动、轻微灯闪或暂时遮挡不得误拆。分屏模板、边框或版式持续存在不等于同一镜头；分屏内容由鞋/腿/手部等局部特写切换为人物面部、上半身或全身，或者上下分屏中的主体和景别整体替换时，必须按真实内容切换新建record。camera_angle只表示摄影机高度方向（平拍/俯拍/仰拍/俯视）；camera_view单独表示水平观察方向（正视/侧视/反打）。每次cut或reframe后都必须重新观察摄影机相对主体的轴线、人物可见正侧面和构图重心，独立填写camera_view及camera_view_evidence，禁止复制上一镜的“正视”。人物自己转头或侧身不能单独证明摄影机变成侧视，必须写直接可见的机位依据。每条record必须单独填写time_of_day=day|night|dawn|dusk|unclear及time_of_day_evidence。先看窗外天空、自然光方向与强弱、日光/夜景可见性、室外环境和明确剧情时间线；室内灯亮不能单独证明夜晚，冷色调或暗曝光也不能单独证明夜晚。证据不足必须写unclear并说明缺少直接昼夜线索。旧字段lighting_composition仅可保留为空字符串兼容，不得再承载生成事实；最终构图只使用generation_facts。
5. narrative_group给出建议剧情镜头组编号，从1连续递增。一个组应覆盖完整的动作或剧情闭环，通常5-15秒；快速插入镜头或场景切换如果共同完成同一个揭示、回忆或反应，可以留在同组。不要按固定秒数拆组。
6. 无法确认的信息使用空字符串并降低confidence，不要把unknown、未知、待确认写进文字字段；只有身份歧义或影响理解的不确定性才写入warnings。
7. 画面人物与声音人物必须分开记录。primary_visible_subject只写该区间画面中心主体；visible_characters必须逐个列出真正出现在画面中的前景、中景和背景人物，包括虚焦、静止、不说话、tiny、只露出半身、贴近画框边缘或被部分遮挡但仍可辨认的人物。先独立清点首、中、尾及动作证据页中出现过的所有人体，再逐一填写人物字段；不能先看资产表再决定保留谁。未绑定资产的人物使用“背景路人A/B/C”，仍完整记录站位与运动，且不得加入asset_hints。场景或物体永远不得进入visible_characters。所有左右方向一律以观众看到的画面为准：frame_left=画面左，frame_right=画面右；禁止用演员自身的左/右，也禁止只写“朝左/朝右”而不标明画面坐标。每个人必须把躯干、头部和眼神分开判断：torso_orientation写正面、背面、朝向镜头、背向镜头、三分之四朝画面左/右、侧身朝画面左/右或unclear；head_orientation独立写头部朝画面左/右、镜头、上/下或组合方向；gaze_direction写眼睛在画面坐标中的方向；gaze_target写具体人物、物件、镜头或画外左/右/上/下，无法确认写unclear。不得因为人物位于另一个人的左/右就推断其朝向对方；只有脸、鼻尖、耳朵、肩线或躯干轴线提供直接证据时才可写面对某人。body_orientation仅作为torso_orientation与head_orientation的简短合并摘要，不得与二者冲突。每个可见人物必须填写：identity、appearance、screen_position、mouth_state、mouth_visual_action、speech_sync_status、mouth_action_evidence，以及 body_posture=standing|sitting|walking|kneeling|crouching|lying|leaning|unclear、torso_orientation、head_orientation、body_orientation、gaze_direction、gaze_target、support_contact、visibility_scope=全身|半身|胸像|局部|背景虚焦、depth_layer=前景|中景|背景、relative_camera_distance=closer_than_focus|same_plane|farther_than_focus|unclear、frame_occupancy=dominant|large|medium|small|tiny、frame_crop、occlusion_relation、depth_evidence、focus_state=sharp|slightly_soft|heavily_defocused|unclear、focus_evidence、facial_expression、character_action。必须先完全忽略清晰度，仅依据人物在画面中的相对尺寸、是否被画框大幅裁切、透视放大和遮挡顺序确定depth_layer与relative_camera_distance；大面积占画、贴边并被裁切的人物通常更靠近镜头，即使其heavily_defocused也仍是前景。然后再独立确定focus_subject与focus_state。浅景深默认只有主焦平面人物为sharp；不同depth_layer同时sharp必须在generation_facts.depth_of_field写deep。镜头级visible_action和expression_gaze只做综合摘要。每条记录还必须填写relative_blocking数组：若画面有N名可见人物，必须对每一对人物恰好填写一次，共N*(N-1)/2项；subject与relative_to必须使用visible_characters中的identity，禁止同一对正反重复。每项填写horizontal_relation=left_of|right_of|overlapping|unclear、depth_relation=in_front_of|behind|same_plane|unclear、occlusion=blocks|blocked_by|none|unclear、facing_relationship=face_to_face|same_direction|back_to_back|crossing|unclear、evidence。facing_relationship必须根据两人的torso_orientation与head_orientation判定；位置左右不能充当朝向证据。不得默认多人same_plane；证据不足写unclear。每条记录还必须填写edge_character_audit={{"status":"present|none|uncertain","observations":[{{"edge":"left|right|top|bottom","identity":"身份或稳定外观标签","visible_fragment":"可见脸部/肩部/服装/身体局部","screen_position":"具体边缘位置"}}]}}；逐一查看四条画框边缘，status=present时每个观察对象都必须进入visible_characters。不得因人物虚焦、无动作、无台词、被裁切、未绑定资产或只露局部而省略。禁止根据台词或画外音推断可见人物。短反应镜头和长台词镜头都必须逐页查看联系图中的区间首部、中部和尾部。
{dialogue_requirement}
9. 输出前逐条复核：{dialogue_qa}再检查候选切点与主体。先把首、中、尾锚点、事件关键帧、运动三联帧和person_trajectory中出现过的所有人体，与visible_characters逐一对账；路人没有资产、只短暂出现或只在背景/边缘可见都不是删除理由。再确认visible_characters中的identity全部是人物，场景、建筑、车辆、喷泉、家具、灯具与物体不得混入。逐个检查visible_characters中的每个人都具备body_posture、torso_orientation、head_orientation、body_orientation、gaze_direction、gaze_target、support_contact、visibility_scope、depth_layer、relative_camera_distance、frame_occupancy、frame_crop、occlusion_relation、depth_evidence、focus_state、focus_evidence、facial_expression、anchor_presence、character_action，不得留空。逐人对照首中尾锚点：中间可见而首尾不可见者，character_action必须解释其短暂出现、经过、遮挡或离开，不得写静止或动作不明。逐人复核躯干朝向、头部朝向、视线方向和视线目标互不替代且不矛盾，所有左右均明确为画面左/右。多人物镜头再复核facing_relationship，禁止把同处一侧、同看画外或前后虚焦构图误写成面对面。复核depth_evidence只能证明物理前后，focus_evidence只能证明光学清晰度；禁止用“虚焦所以在后方”或“清晰所以在前景”的逻辑。检查relative_blocking覆盖全部人物配对、无反向重复、无前后遮挡矛盾，并让spatial_positions按relative_blocking从近到远总结；检查edge_character_audit四边无漏人。每次cut或reframe后复核camera_view与camera_view_evidence，确认没有机械沿用上一镜视角。再做一次“动作字段→道具表”反查：逐条扫描所有人物动作、visible_action、contact_actions、initial_frame和final_frame，凡出现持有、携带、递接、开合、展示或使用的物件，prop_continuity必须存在同名记录；再按时间顺序检查同一道具的end_state与下一record的start_state、holder和source是否相容。已在前镜出现的道具，后镜不得标成首次揭示；证据不足写unclear和warning，不能留空。最后复核generation_facts：主体受光、明暗、昼夜与景深可以保留；固定灯带、吊灯、顶光、霓虹和彩色氛围灯以及静态蓝光/红光必须删除，只有开启/熄灭/闪烁/扫过/屏幕光照脸等剧情性变化可保留。每个字段写完整事实短句，本范围允许使用约24000个有效输出token；不得用重复套话凑字数。仅返回JSON对象，不要Markdown。

根对象必须包含dialogue_events和boundary_audit。boundary_audit结构为[{{"candidate_seconds":0.433,"decision":"keep_cut|keep_reframe|merge_false","observed_seconds":0.433,"before_visible_subject":"切点前主画面主体","after_visible_subject":"切点后主画面主体","same_shot_evidence":false,"reason":"简短可见依据"}}]，并覆盖每个内部候选切点。before/after标签必须与相邻records的primary_visible_subject一致；locked候选若保留必须使用原时间，若merge_false必须有直接同镜证据。
下面JSON中的0.0和1.2仅为字段结构示意；实际第一条必须使用本次range_start，所有时间均为原视频绝对秒数。除示例字段外，根对象必须增加dialogue_events数组；每条record必须增加interval_id、evidence_interval_ids、start_anchor_subject、middle_anchor_subject、end_anchor_subject、anchor_evidence_timestamps、dialogue_event_ids、dialogue_placements；每个dialogue_utterances元素必须增加dialogue_id。JSON必须是：
{{"schema_version":"1.0","episode":"{episode}","duration_seconds":{duration:.3f},"records":[{{"id":"E001","narrative_group":1,"beat_summary":"本组剧情动作闭环","start_seconds":0.0,"end_seconds":1.2,"boundary":"cut|reframe|continuous","transition":"进入与离开本镜头的具体方式","scene_observation":"场景识别摘要","fixed_scene_evidence":"仅供资产匹配的固定装修与常设陈设","primary_visible_subject":"本区间主画面人物或物件","focus_subject":"实际最清晰的主体","focus_policy":"locked|rack_focus|unclear","visible_characters":[{{"identity":"中文身份、资产人物名或背景路人A/B/C；只能是人","appearance":"可见服装与外形","screen_position":"画面位置","mouth_state":"speaking|not_speaking|unclear","mouth_visual_action":"闭合/短暂张开后闭合/连续开合/微笑露齿/抿嘴/unclear等可见事实","speech_sync_status":"matched_on_screen_dialogue|off_screen_audio_reaction|non_speech_mouth_motion|no_visible_mouth_motion|unclear","mouth_action_evidence":"时间点或首中尾口型变化证据","body_posture":"standing|sitting|walking|kneeling|crouching|lying|leaning|unclear","torso_orientation":"正面/背面/朝向镜头/背向镜头/三分之四朝画面左或右/侧身朝画面左或右/unclear","head_orientation":"头部朝画面左/右/镜头/上/下或组合方向/unclear","body_orientation":"躯干与头部朝向的简短合并摘要","gaze_direction":"眼睛相对画面的方向或unclear","gaze_target":"具体人物/物件/镜头/画外左右上下或unclear","support_contact":"站立着地、坐在椅子、倚靠或被搀扶等","visibility_scope":"全身|半身|胸像|局部|背景虚焦","depth_layer":"前景|中景|背景","relative_camera_distance":"closer_than_focus|same_plane|farther_than_focus|unclear","frame_occupancy":"dominant|large|medium|small|tiny","frame_crop":"画框如何裁切身体及贴哪侧边缘","occlusion_relation":"遮挡谁、被谁遮挡或none","depth_evidence":"人物尺寸、透视、裁切和遮挡形成的物理距离依据","focus_state":"sharp|slightly_soft|heavily_defocused|unclear","focus_evidence":"仅写眼睛、五官、发丝与轮廓边缘的光学清晰或柔化依据；不得写衣服或配饰","facial_expression":"该人物自己的可见表情或unclear","character_action":"该人物自己的可见动作；无显著动作写静止状态"}}],"relative_blocking":[{{"subject":"人物A","relative_to":"人物B","horizontal_relation":"left_of|right_of|overlapping|unclear","depth_relation":"in_front_of|behind|same_plane|unclear","occlusion":"blocks|blocked_by|none|unclear","facing_relationship":"face_to_face|same_direction|back_to_back|crossing|unclear","evidence":"人物尺寸、裁切、透视、遮挡、躯干与头部朝向依据"}}],"edge_character_audit":{{"status":"present|none|uncertain","observations":[{{"edge":"left|right|top|bottom","identity":"身份或稳定外观标签","visible_fragment":"脸部/肩部/身体局部；已绑定人物不得复述服装","screen_position":"具体边缘位置"}}]}},"visual_details":"仅供资产识别的原始视觉证据","transient_visual_details":"临时剧情物件、状态与位置","generation_facts":{{"subject_lighting":"主体受光正常，或具体异常受光","contrast_exposure":"仅写整体明暗、曝光与反差","depth_of_field":"shallow|deep|rack_focus|unclear","focus_transition":"无变化或具体焦点转移","composition_change":"构图重心、前景遮挡或景别变化","dynamic_environment":"门窗/车辆/烟雨/剧情性灯光变化；无则写无"}},"lighting_composition":"","visible_action":"全体人物动作综合摘要","shot_size":"wide|full|medium|close-up|extreme-close-up|unknown","camera_angle":"eye-level|high-angle|low-angle|overhead|unknown","camera_motion":"static|pan|tilt|push-in|pull-out|tracking|handheld|unknown","camera_view":"1 俯瞰/俯视|2 正视|3 侧视|4 反打|待确认","spatial_positions":"按relative_blocking与镜头距离从近到远写人物占比、裁切、遮挡、画面左右与朝向关系","contact_actions":"精确的人物—人物或人物—物体接触动作链；无则写无","initial_frame":"仅写人物/物体位置动作和瞬时状态","final_frame":"仅写人物/物体位置动作和瞬时状态","expression_gaze":"全体人物表情与视线综合摘要","mouth_dynamics":"逐人物汇总嘴部可见动作与台词同步状态，不据单帧张嘴推断发言","speaker":"声音说话人的中文身份或空字符串","dialogue_utterances":[{{"start_seconds":0.2,"end_seconds":0.7,"speaker":"中文身份","speaker_visibility":"on_screen|off_screen|unclear","text":"逐句中文原台词","basis":"audio|burned_subtitle|both"}}],"chinese_dialogue":"按时间顺序用；连接全部逐句台词，或空字符串","on_screen_text":"文字类型与内容或空字符串","sound":"声音或空字符串","asset_hints":["中文源资产名"],"confidence":"high|medium|low","warnings":[]}}],"warnings":[],"errors":[]}}
"""
    rendered = rendered.replace(
        '"appearance":"可见服装与外形"',
        '"appearance":"匹配人物-服装态时固定写asset_bound；未匹配时仅写最小识别依据"',
        1,
    )
    rendered = rendered.replace(
        '"visual_details":"颜色、材质、物件、前中后景及位置"',
        '"visual_details":"仅供第一阶段资产识别的场景与物件事实","transient_visual_details":"不属于已绑定人物服装态或场景资产的临时剧情物件、状态与位置"',
        1,
    )
    rendered = rendered.replace(
        '"fixed_scene_evidence":"仅供资产匹配的固定装修与常设陈设",',
        '"fixed_scene_evidence":"仅供资产匹配的固定装修与常设陈设","time_of_day":"day|night|dawn|dusk|unclear","time_of_day_evidence":"窗外天空、自然光、室外环境或明确剧情时间线依据；不以室内灯光单独判断",',
        1,
    )
    rendered = rendered.replace(
        '"transient_visual_details":"临时剧情物件、状态与位置","generation_facts":',
        '"transient_visual_details":"临时剧情物件、状态与位置","prop_continuity":[{"instance_id":"本范围稳定P01/P02","visible_instance_count":1,"identity":"稳定道具名","appearance_invariants":"区分同一实例的形状、尺寸、结构；翻面颜色变化不新建实例","holder":"人物名或none","support_contact":"具体手部/身体/台面接触","start_state":"首帧可见位置、开合与持有状态","middle_action":"中间连续物理动作；无则写保持","end_state":"尾帧可见位置、开合与持有状态","source":"上一镜继承/口袋/包/抽屉/画外递入/硬切后首次揭示/unclear","continuity_from_previous":"继承说明或none","continuity_to_next":"下一镜应保持的状态或none","exclusive_holder_after":"交接完成后的唯一持有者；无交接写当前持有者或none"}],"generation_facts":',
        1,
    )
    rendered = rendered.replace(
        '"focus_policy":"locked|rack_focus|unclear","visible_characters"',
        '"focus_policy":"locked|rack_focus|unclear","screen_order_left_to_right":{"start":["人物A","人物B"],"middle":["人物A","人物B"],"end":["人物B","人物A"]},"position_transition":{"type":"stable|crosses_behind|crosses_in_front|reblocked_after_cut|unclear","description":"按观众画面坐标描述可见换位路径或硬切后重新构图","evidence_timestamps":[0.5,1.0]},"visible_characters"',
        1,
    )
    rendered = rendered.replace(
        '"facial_expression":"该人物自己的可见表情或unclear","character_action"',
        '"facial_expression":"该人物自己的可见表情或unclear","anchor_presence":{"start":true,"middle":true,"end":true},"character_action"',
        1,
    )
    return rendered.replace(
        '"camera_view":"1 俯瞰/俯视|2 正视|3 侧视|4 反打|待确认","spatial_positions"',
        '"camera_view":"1 俯瞰/俯视|2 正视|3 侧视|4 反打|待确认","camera_view_evidence":"摄影机相对主体轴线、人物可见正侧面与构图重心的直接依据","spatial_positions"',
        1,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description="Ask 中转站 for a compact full-episode evidence timeline from a temporary TOS URL.")
    parser.add_argument("--url-file", type=Path, required=True)
    parser.add_argument("--key-pool-file", type=Path, required=True)
    parser.add_argument("--key-offset", type=int, default=0, help="Rotate the key pool so concurrent workers start on different keys")
    parser.add_argument("--single-key-only", action="store_true", help="Use only the leased key selected by --key-offset; required under cross-window relay slot leasing")
    parser.add_argument("--episode", required=True)
    parser.add_argument("--duration", type=float, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--raw-output", type=Path, required=True)
    parser.add_argument("--boundaries-file", type=Path)
    parser.add_argument("--context-sheet", type=Path, action="append", default=[], help="Repeat for the range edge sheet and each chronological contact-sheet page")
    parser.add_argument("--evidence-pack", type=Path, help="Role-labelled adaptive evidence manifest; cannot be combined with anonymous context or pose sheets")
    parser.add_argument("--pose-evidence-sheet", type=Path, help="Optional JointBDOE + per-crop MediaPipe Pose Lite geometry page")
    parser.add_argument("--assets-workbook", type=Path, help="Inject all episode-filtered person, scene, and prop reference cards from the asset workbook")
    parser.add_argument("--dialogue-injection", type=Path, required=True, help="Validated dialogue injection built from the current episode ledger")
    parser.add_argument("--dialogue-ledger", type=Path, required=True, help="Current validated whole-episode ledger used to verify the injection fingerprint")
    parser.add_argument("--continuity-injection", type=Path, help="Previous logical core tail state for seam continuity; observation context only")
    parser.add_argument("--clip-manifest", type=Path, help="Overlap-clip manifest; enables clip-to-absolute timestamp normalization")
    parser.add_argument("--range-id", help="Range id in --clip-manifest")
    parser.add_argument("--range-start", type=float, default=0.0)
    parser.add_argument("--range-end", type=float)
    parser.add_argument("--model", default="gemini-3.6-flash")
    parser.add_argument("--endpoint", default="https://api.b.ai/v1/chat/completions")
    parser.add_argument("--timeout", type=int, default=1000)
    parser.add_argument("--max-output-tokens", type=int, default=30000)
    args = parser.parse_args()

    if args.evidence_pack and (args.context_sheet or args.pose_evidence_sheet):
        raise RuntimeError("--evidence-pack cannot be combined with --context-sheet or --pose-evidence-sheet")

    url_data = json.loads(args.url_file.read_text(encoding="utf-8-sig"))
    signed_url = url_data.get("signed_url")
    if not signed_url:
        raise RuntimeError("Temporary URL file has no signed_url")
    keys = read_keys(args.key_pool_file.resolve())
    key_offset = args.key_offset % len(keys)
    keys = keys[key_offset:] + keys[:key_offset]
    if args.single_key_only:
        keys = keys[:1]
    boundary_data: dict[str, Any] = {}
    if args.boundaries_file:
        boundary_data = json.loads(args.boundaries_file.read_text(encoding="utf-8-sig"))
    range_end = args.duration if args.range_end is None else args.range_end
    if args.range_start < 0 or range_end > args.duration or range_end <= args.range_start:
        raise RuntimeError("Invalid analysis range")
    clip_mapping = None
    if args.clip_manifest:
        if not args.range_id:
            raise RuntimeError("--range-id is required with --clip-manifest")
        clip_mapping = load_clip_mapping(args.clip_manifest.resolve(), args.range_id)
        if abs(float(clip_mapping["logical_start_seconds"]) - args.range_start) > 0.001 or abs(float(clip_mapping["logical_end_seconds"]) - range_end) > 0.001:
            raise RuntimeError("clip mapping logical range does not match --range-start/--range-end")
    asset_cards = episode_asset_reference_cards(args.assets_workbook.resolve(), str(args.episode).zfill(2)) if args.assets_workbook else []
    dialogue_injection = load_dialogue_injection(
        args.dialogue_injection.resolve(),
        args.dialogue_ledger.resolve(),
        str(args.episode).zfill(2),
        args.range_start,
        range_end,
    )
    continuity_injection: dict[str, Any] = {}
    if args.continuity_injection:
        continuity_injection = json.loads(args.continuity_injection.resolve().read_text(encoding="utf-8-sig"))
        if continuity_injection.get("schema_version") != "1.0-seam-continuity-injection":
            raise RuntimeError("continuity injection must use schema_version 1.0-seam-continuity-injection")
    evidence_pages: list[dict[str, Any]] = []
    evidence_manifest: dict[str, Any] = {}
    if args.evidence_pack:
        evidence_pages, evidence_manifest = load_evidence_pack(args.evidence_pack)
    evidence_roles = [item["role"] for item in evidence_pages]
    has_pose_evidence = "multi_person_geometry" in evidence_roles or bool(args.pose_evidence_sheet)
    has_context_evidence = bool(evidence_pages or args.context_sheet)
    content = [{"type": "image_url", "image_url": {"url": signed_url}}]
    if clip_mapping:
        content.append({"type": "text", "text": clip_transport_note(clip_mapping)})
    for page in evidence_pages:
        content.append({"type": "text", "text": f"PAGE_ROLE={page['role']}. Absolute-time local evidence; clipped video and unmarked original pixels have higher priority than overlays."})
        if page["role"] == "person_trajectory" and page.get("summary"):
            content.append({"type": "text", "text": "TRAJECTORY_GEOMETRY_SUMMARY=" + json.dumps(page["summary"], ensure_ascii=False, separators=(",", ":"))})
        content.append({"type": "image_url", "image_url": {"url": image_data_url(page["path"])}})
    for context_sheet in args.context_sheet:
        if not context_sheet.exists():
            raise RuntimeError(f"Context sheet does not exist: {context_sheet}")
        content.append({"type": "image_url", "image_url": {"url": image_data_url(context_sheet)}})
    if args.pose_evidence_sheet:
        if not args.pose_evidence_sheet.exists():
            raise RuntimeError(f"Pose evidence sheet does not exist: {args.pose_evidence_sheet}")
        content.append({"type": "text", "text": "PAGE_ROLE=multi_person_geometry. Auxiliary local geometry evidence; original video and unmarked frames have higher priority."})
        content.append({"type": "image_url", "image_url": {"url": image_data_url(args.pose_evidence_sheet)}})
    boundary_contract = build_boundary_contract(boundary_data, args.range_start, range_end)
    content.append({"type": "text", "text": prompt(str(args.episode).zfill(2), args.duration, boundary_contract, args.range_start, range_end, has_context_evidence, has_pose_evidence, evidence_roles, asset_cards, dialogue_injection, continuity_injection)})
    payload = {
        "model": args.model,
        "messages": [{"role": "user", "content": content}],
        "stream": False,
        "max_tokens": args.max_output_tokens,
        "temperature": 0.1,
        "response_format": {"type": "json_object"},
    }

    retryable = {429, 500, 502, 503, 504, 520, 522, 524}
    last_error = ""
    data: dict[str, Any] | None = None
    attempts = 0
    # A quota failure belongs to one relay key, not to the whole pool. Try every
    # configured key before failing so concurrent ranges do not stop on a
    # depleted first-choice key.
    max_attempts = 2 if args.single_key_only else max(2, len(keys))
    for attempts in range(1, max_attempts + 1):
        response = requests.post(
            args.endpoint,
            headers={"Authorization": f"Bearer {keys[(attempts - 1) % len(keys)]}", "Content-Type": "application/json; charset=utf-8"},
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            timeout=(60, args.timeout),
        )
        if response.status_code < 400:
            data = response.json()
            break
        last_error = f"中转站 HTTP {response.status_code}: {response.text[:2000]}"
        quota_exhausted = response.status_code == 400 and (
            "insufficient_user_quota" in response.text
            or "credit insufficient balance" in response.text.lower()
        )
        unusable_key = response.status_code in {401, 403} and (
            "api_key" in response.text.lower()
            or "unauthorized" in response.text.lower()
            or "鉴权" in response.text
            or "insufficient" in response.text.lower()
        )
        if args.single_key_only and (quota_exhausted or unusable_key):
            raise RuntimeError(last_error)
        if (response.status_code not in retryable and not quota_exhausted and not unusable_key) or attempts == max_attempts:
            raise RuntimeError(last_error)
        # Quota errors need key rotation, not a long backoff.
        if not quota_exhausted and not unusable_key:
            time.sleep(5 * attempts + random.random())
    if data is None:
        raise RuntimeError(last_error or "中转站 request failed")

    text = response_text(data)
    raw_record = {
        "provider": "bai-tos-url",
        "prompt_contract_version": PROMPT_CONTRACT_VERSION,
        "model": data.get("model") or args.model,
        "request_id": data.get("id", ""),
        "usage": data.get("usage", {}),
        "attempts": attempts,
        "key_slot": key_offset + 1,
        "max_output_tokens": args.max_output_tokens,
        "context_sheet": has_context_evidence,
        "context_sheet_count": len(evidence_pages) + len(args.context_sheet),
        "evidence_pack": args.evidence_pack.name if args.evidence_pack else None,
        "evidence_roles": evidence_roles,
        "pose_evidence_sheet": has_pose_evidence,
        "pose_gate": evidence_manifest.get("pose_gate") if evidence_manifest else None,
        "episode_asset_reference_count": len(asset_cards),
        "episode_asset_reference_counts": {
            asset_type: sum(card.startswith(f"- [{asset_type}]") for card in asset_cards)
            for asset_type in ("人物", "场景", "道具")
        },
        "dialogue_injection": True,
        "dialogue_injection_event_count": len(dialogue_injection.get("analysis_events") or []),
        "continuity_injection": bool(continuity_injection),
        "transport_mode": "overlap_clipped_video" if clip_mapping else "full_episode_video",
        "range_id": args.range_id if clip_mapping else None,
        "text": text,
    }
    atomic_json(args.raw_output.resolve(), raw_record)
    repaired = False
    try:
        parsed = extract_json(text)
    except (json.JSONDecodeError, RuntimeError) as parse_error:
        repair_payload = {
            "model": args.model,
            "messages": [{"role": "user", "content": "修复下面损坏的JSON。只纠正引号、逗号、括号和截断的结构，不改写事实、时间、字段或文本；只返回一个有效JSON对象。\n\n" + text}],
            "stream": False,
            "max_tokens": args.max_output_tokens,
            "temperature": 0,
            "response_format": {"type": "json_object"},
        }
        repair_response = requests.post(
            args.endpoint,
            headers={"Authorization": f"Bearer {keys[0]}", "Content-Type": "application/json; charset=utf-8"},
            data=json.dumps(repair_payload, ensure_ascii=False).encode("utf-8"),
            timeout=(60, args.timeout),
        )
        if repair_response.status_code >= 400:
            raise RuntimeError(f"中转站 JSON repair HTTP {repair_response.status_code}: {repair_response.text[:2000]}") from parse_error
        repair_data = repair_response.json()
        repair_text = response_text(repair_data)
        parsed = extract_json(repair_text)
        repaired = True
        raw_record["initial_parse_error"] = str(parse_error)
        raw_record["repair"] = {
            "model": repair_data.get("model") or args.model,
            "request_id": repair_data.get("id", ""),
            "usage": repair_data.get("usage", {}),
            "text": repair_text,
        }
        atomic_json(args.raw_output.resolve(), raw_record)
    parsed.setdefault("schema_version", "1.0")
    parsed["prompt_contract_version"] = PROMPT_CONTRACT_VERSION
    parsed.setdefault("episode", str(args.episode).zfill(2))
    parsed.setdefault("duration_seconds", args.duration)
    parsed.setdefault("analysis_start_seconds", args.range_start)
    parsed.setdefault("analysis_end_seconds", range_end)
    parsed.setdefault("warnings", [])
    parsed.setdefault("errors", [])
    timestamp_mode = normalize_clipped_timestamps(parsed, clip_mapping)
    parsed["transport"] = {
        "mode": "overlap_clipped_video" if clip_mapping else "full_episode_video",
        "range_id": args.range_id if clip_mapping else None,
        "timestamp_normalization": timestamp_mode,
        "clip_start_seconds": float(clip_mapping["clip_start_seconds"]) if clip_mapping else 0.0,
    }
    parsed["boundary_contract"] = boundary_contract
    for record in parsed.get("records") or []:
        if not record.get("visible_action") and record.get("visual_action"):
            record["visible_action"] = record.pop("visual_action")
    hydrate_dialogue_injection(parsed, dialogue_injection)
    atomic_json(args.output.resolve(), parsed)
    print(json.dumps({"status": "complete", "model": data.get("model") or args.model, "records": len(parsed.get("records") or []), "attempts": attempts, "repaired": repaired}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
